"""审核项 F05/F06/F07 回归（契约 p3-imports-analysis.md §1.1/§1.2 v2.1）。

- F05 整批健康检查：合法 + 损坏文件混批 → confirm 409 列出失败文件、
  零业务写入、batch 保持 pending；移除损坏文件重新 preview 后成功。
- F06 身份候选：历史学年同号同名 → preview 出 identity_candidates
  （known 不计）；confirm 无确认 → 409 附完整候选清单（零写入）；
  带identity_confirmations → 接续旧 identity（补登本学年 alias）；
  确认伪造 person → 422；跨届同号不同名 → 撞号整文件拒绝。
  未标学年（NULL）的 alias 登记不再默认视为本学年身份（契约 v2.2/G06，
  旧例外违反永久身份红线）：无有效期证据时进候选，仅当该行自身
  valid_from/valid_to 明确覆盖目标文件日期才直接命中 known。
- F07 成员有效期：新建/接续成员 valid_from = max(学年 start_date,
  文件日期)，绝不用"今天"；确认接续的既有人 homeroom 补建 Enrollment、
  teaching 补 TeachingClassMember（幂等），读侧名册可见。
- 考试日期冲突（v2.1 边界裁决）：批内或与库内同 exam_name 不同
  exam_date → preview 计 warnings + confirm 整批 409。

造表 helper 与 test_p3_imports_e01_e02_e05.py 同源复制（conftest 不在
本任务允许清单）。
"""

from datetime import date
from pathlib import Path

import pytest

API = "/api/v1"

EXAM_F5 = "P3F5X"
EXAM_F6 = "P3F6X"
EXAM_F7A = "P3F7A"
EXAM_F7B = "P3F7B"
EXAM_F7T = "P3F7T"

_HEADERS = {
    1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
    5: "语文", 6: "数学", 7: "英语",
    8: "物理", 9: "物理等级分",
    20: "+3总分", 21: "主三门总分", 22: "主三门百分位", 23: "主三门学籍排名",
    24: "3+3总分", 25: "3+3百分位", 26: "3+3学籍排名",
    27: "语文百分位", 28: "数学百分位", 29: "英语百分位",
}


def _student_row(
    sid, name, *, class_num=6, chinese=None, math=None, english=None,
    physics=None, physics_grade=None, teaching_label=None,
):
    data = {1: sid, 2: class_num, 3: 1, 4: name}
    if chinese is not None:
        data[5] = chinese
    if math is not None:
        data[6] = math
    if english is not None:
        data[7] = english
    if physics is not None:
        data[8] = physics
    if physics_grade is not None:
        data[9] = physics_grade
    if teaching_label is not None:
        data[30] = teaching_label
    return data


def _write_grade23_xlsx(path: Path, rows, teaching_col: bool = False) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    headers = dict(_HEADERS)
    if teaching_col:
        headers[30] = "教学班"
    for col, text in headers.items():
        ws.cell(row=2, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def _upload(client, paths, mode: str, extra: dict | None = None):
    """multipart preview：支持单文件或文件列表（F05 混批需要多文件）。"""
    if isinstance(paths, (str, Path)):
        paths = [paths]
    data = {"mode": mode}
    if extra:
        data.update({k: str(v) for k, v in extra.items()})
    files = []
    for p in paths:
        with open(p, "rb") as fh:
            files.append(
                (
                    "files",
                    (
                        p.name,
                        fh.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                )
            )
    return client.post(f"{API}/imports/preview", data=data, files=files)


def _db():
    from app.db.models import SessionLocal

    return SessionLocal()


def _make_history_person(db, year_name: str, start: str, end: str, domain: str,
                         display_name: str, alias_value: str):
    """造一个"历史学年的人"：独立学年 + 域内 identity + 标了该学年的 alias。

    返回 (identity_id, academic_year_id)。alias 明确标注非本学年的
    academic_year_id——这正是 F06 要求走候选确认流程（而非自动接续）
    的场景。"""
    from app.db import workspace_models as wm

    ay = wm.AcademicYear(
        name=year_name,
        start_date=date.fromisoformat(start),
        end_date=date.fromisoformat(end),
    )
    db.add(ay)
    db.flush()
    ident = wm.WsStudentIdentity(data_domain=domain, display_name=display_name)
    db.add(ident)
    db.flush()
    db.add(
        wm.WsStudentAlias(
            identity_id=ident.id,
            alias_value=alias_value,
            data_domain=domain,
            academic_year_id=ay.id,
            alias_scope=str(ay.id),
            source="synthetic-test",
        )
    )
    db.commit()
    return ident.id, ay.id


# ────────────────────────────── F05 整批健康检查 ──────────────────────────────


def test_f05_mixed_batch_rejected_all_or_nothing(client, v1_seed, tmp_path):
    """合法 + 损坏文件混批：preview parsed_ok=[true,false]；confirm 409
    列出失败文件、ScoreFact 零残留、batch 保持 pending；移除损坏文件
    重新 preview → confirm 成功（绝不静默部分入库）。"""
    from app.db import workspace_models as wm

    good = _write_grade23_xlsx(
        tmp_path / "F05好表.xlsx",
        [_student_row("2025H6-55", "秦伍", chinese=80, math=85)],
    )
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"this is not a xlsx file")

    prev = _upload(
        client, [good, broken], "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_F5, "exam_date": "2025-11-10"},
    )
    assert prev.status_code == 200, prev.text
    items = prev.json()["items"]
    assert [i["parsed_ok"] for i in items] == [True, False]
    assert items[1]["kind"] == "unknown"
    assert items[1]["message"]

    r = client.post(f"{API}/imports/confirm", json={"token": prev.json()["token"]})
    assert r.status_code == 409, r.text
    body = r.json()
    assert body["error"] == "link_version_conflict"
    failed = body["failed_files"]
    assert [f["filename"] for f in failed] == ["broken.xlsx"]
    assert failed[0]["kind"] == "unknown"

    db = _db()
    try:
        # 零业务写入：好文件的成绩也不入库（E05 整批原子）
        assert db.query(wm.ScoreFact).filter_by(exam_name=EXAM_F5).count() == 0
        batch = db.query(wm.ImportBatch).filter_by(token=prev.json()["token"]).one()
        assert batch.status == "pending"
    finally:
        db.close()

    # 移除损坏文件重新 preview → confirm 成功
    prev2 = _upload(
        client, good, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_F5, "exam_date": "2025-11-10"},
    )
    assert prev2.status_code == 200
    ok = client.post(
        f"{API}/imports/confirm", json={"token": prev2.json()["token"]}
    )
    assert ok.status_code == 200, ok.text
    assert ok.json()["imported"] == 2  # 语文 + 数学


# ────────────────────────────── F06 身份候选 ──────────────────────────────


def test_f06_history_alias_candidate_flow(client, v1_seed, tmp_path):
    """旧学年同号同名：preview 出候选（known 不含）→ 无确认 409 附候选
    零写入 → 显式确认后接续旧 identity、补登本学年 alias、补建学籍；
    再 preview 同文件 → 本学年命中计入 known。"""
    from app.db import workspace_models as wm

    db = _db()
    try:
        old_pid, old_ay_id = _make_history_person(
            db, "2024-2025", "2024-09-01", "2025-07-15",
            "homeroom", "秦旧", "2024H6-88",
        )
    finally:
        db.close()

    path = _write_grade23_xlsx(
        tmp_path / "F06接续.xlsx",
        [_student_row("2024H6-88", "秦旧", chinese=77, math=88)],
    )
    prev = _upload(
        client, path, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_F6, "exam_date": "2025-11-10"},
    )
    assert prev.status_code == 200, prev.text
    item = prev.json()["items"][0]
    # 历史命中：不算 known、不算 new，只出候选
    assert item["known_students"] == 0
    assert item["new_students"] == []
    assert item["identity_candidates"] == [
        {
            "alias": "2024H6-88",
            "name": "秦旧",
            "person_id": old_pid,
            "academic_year_id": old_ay_id,
            "academic_year_name": "2024-2025",
            "basis": "history_alias",
        }
    ]
    assert any("历史学年" in w for w in item["warnings"])

    # 未逐项确认 → 409 + 完整候选清单，零写入，batch 保持 pending
    denied = client.post(
        f"{API}/imports/confirm", json={"token": prev.json()["token"]}
    )
    assert denied.status_code == 409, denied.text
    body = denied.json()
    assert body["error"] == "link_version_conflict"
    assert body["identity_candidates"] == item["identity_candidates"]

    db = _db()
    try:
        assert db.query(wm.ScoreFact).filter_by(exam_name=EXAM_F6).count() == 0
        assert (
            db.query(wm.Enrollment)
            .filter_by(admin_class_id=v1_seed.h6_id, identity_id=old_pid)
            .count()
            == 0
        )
        batch = db.query(wm.ImportBatch).filter_by(token=prev.json()["token"]).one()
        assert batch.status == "pending"
    finally:
        db.close()

    # 同 token 显式确认 → 接续旧 identity
    ok = client.post(
        f"{API}/imports/confirm",
        json={
            "token": prev.json()["token"],
            "identity_confirmations": {"2024H6-88": old_pid},
        },
    )
    assert ok.status_code == 200, ok.text
    result = ok.json()
    assert result["imported"] == 2
    assert result["students_created"] == 0  # 接续既有人，不是新建
    assert result["members_synced"] == 1  # F07：补建目标行政班学籍

    db = _db()
    try:
        fact = (
            db.query(wm.ScoreFact)
            .filter_by(exam_name=EXAM_F6, subject="语文")
            .one()
        )
        assert fact.identity_id == old_pid  # 成绩落到历史 identity
        alias_now = (
            db.query(wm.WsStudentAlias)
            .filter_by(
                alias_value="2024H6-88",
                data_domain="homeroom",
                academic_year_id=v1_seed.ay_id,
            )
            .one()
        )
        assert alias_now.identity_id == old_pid  # 本学年 alias 已补登
        enrollment = (
            db.query(wm.Enrollment)
            .filter_by(admin_class_id=v1_seed.h6_id, identity_id=old_pid)
            .one()
        )
        # F07：有效期 = max(学年 start 2025-09-01, 文件日期 2025-11-10)
        assert enrollment.valid_from == date(2025, 11, 10)
        assert enrollment.status == "active"
    finally:
        db.close()

    # 读侧名册（新学年上下文）能看到该成员："导入成功却看不到"已消除
    roster = client.get(
        f"{API}/homeroom/students",
        params={"academic_year_id": v1_seed.ay_id, "class_id": v1_seed.h6_id},
    )
    assert roster.status_code == 200, roster.text
    assert old_pid in {s["person_id"] for s in roster.json()["students"]}

    # 再 preview 同文件：本学年 alias 已补登 → 直接命中 known，无候选
    again = _upload(
        client, path, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": "P3F6Y", "exam_date": "2025-12-20"},
    )
    assert again.status_code == 200
    item2 = again.json()["items"][0]
    assert item2["known_students"] == 1
    assert item2["identity_candidates"] == []


def test_f06_forged_and_cross_cohort_name_mismatch(client, v1_seed, tmp_path):
    """确认伪造 person（未持有该 alias 历史登记）→ 422；跨届同号不同名
    → 撞号整文件 409 拒绝，不自动改名合并。"""
    from app.db import workspace_models as wm

    db = _db()
    try:
        old_pid, _old_ay_id = _make_history_person(
            db, "2023-2024", "2023-09-01", "2024-07-15",
            "homeroom", "秦古", "2023H6-88",
        )
    finally:
        db.close()

    path = _write_grade23_xlsx(
        tmp_path / "F06伪造.xlsx",
        [_student_row("2023H6-88", "秦古", chinese=66, math=77)],
    )
    prev = _upload(
        client, path, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": "P3F6Z", "exam_date": "2025-11-12"},
    )
    assert prev.status_code == 200
    token = prev.json()["token"]

    # 确认到别人（乙从未持有 2023H6-88 的历史登记）→ 422
    forged = client.post(
        f"{API}/imports/confirm",
        json={
            "token": token,
            "identity_confirmations": {"2023H6-88": v1_seed.yi_h_id},
        },
    )
    assert forged.status_code == 422, forged.text
    assert forged.json()["error"] == "invalid_scope_param"

    db = _db()
    try:
        assert db.query(wm.ScoreFact).filter_by(exam_name="P3F6Z").count() == 0
        batch = db.query(wm.ImportBatch).filter_by(token=token).one()
        assert batch.status == "pending"
    finally:
        db.close()

    # 跨届同号不同名：撞号防呆 → confirm 整文件 409（优先于候选确认）
    bad = _write_grade23_xlsx(
        tmp_path / "F06撞号.xlsx",
        [_student_row("2023H6-88", "秦别人", chinese=66, math=77)],
    )
    prev_bad = _upload(
        client, bad, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": "P3F6W", "exam_date": "2025-11-13"},
    )
    assert prev_bad.status_code == 200
    item = prev_bad.json()["items"][0]
    assert any("同学号不同姓名" in w for w in item["warnings"])

    r = client.post(
        f"{API}/imports/confirm", json={"token": prev_bad.json()["token"]}
    )
    assert r.status_code == 409
    body = r.json()
    assert body["error"] == "link_version_conflict"
    assert body["conflicts"]  # 撞号冲突，而非候选清单

    db = _db()
    try:
        assert db.query(wm.ScoreFact).filter_by(exam_name="P3F6W").count() == 0
    finally:
        db.close()


# ────────────────────────────── F07 成员有效期 ──────────────────────────────


def test_f07_new_member_valid_from_max_rule(client, v1_seed, tmp_path):
    """新建成员 valid_from = max(学年 start_date, 文件日期)：
    文件日期晚于学年 start → 取文件日期；早于学年 start → 取学年 start。
    绝不取"今天"（运行日 2026-09-11 ≠ 断言值）。"""
    from app.db import workspace_models as wm

    # ① 文件日期 2025-11-10 晚于学年 start 2025-09-01 → 取文件日期
    path_a = _write_grade23_xlsx(
        tmp_path / "F7晚.xlsx",
        [_student_row("2025H6-71", "秦柒", chinese=71, math=72)],
    )
    prev_a = _upload(
        client, path_a, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_F7A, "exam_date": "2025-11-10"},
    )
    assert prev_a.status_code == 200
    ok_a = client.post(
        f"{API}/imports/confirm", json={"token": prev_a.json()["token"]}
    )
    assert ok_a.status_code == 200, ok_a.text

    # ② 文件日期 2025-08-20 早于学年 start → 取学年 start 2025-09-01
    path_b = _write_grade23_xlsx(
        tmp_path / "F7早.xlsx",
        [_student_row("2025H6-72", "秦捌", chinese=81, math=82)],
    )
    prev_b = _upload(
        client, path_b, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_F7B, "exam_date": "2025-08-20"},
    )
    assert prev_b.status_code == 200
    ok_b = client.post(
        f"{API}/imports/confirm", json={"token": prev_b.json()["token"]}
    )
    assert ok_b.status_code == 200, ok_b.text

    db = _db()
    try:
        qi7 = db.query(wm.WsStudentIdentity).filter_by(display_name="秦柒").one()
        enrollment7 = (
            db.query(wm.Enrollment)
            .filter_by(admin_class_id=v1_seed.h6_id, identity_id=qi7.id)
            .one()
        )
        assert enrollment7.valid_from == date(2025, 11, 10)  # 文件日期较晚者
        qi8 = db.query(wm.WsStudentIdentity).filter_by(display_name="秦捌").one()
        enrollment8 = (
            db.query(wm.Enrollment)
            .filter_by(admin_class_id=v1_seed.h6_id, identity_id=qi8.id)
            .one()
        )
        assert enrollment8.valid_from == date(2025, 9, 1)  # 学年 start 较晚者
    finally:
        db.close()


def test_f07_teaching_continuation_member_synced(client, v1_seed, tmp_path):
    """teaching 模式确认接续既有人：补 TeachingClassMember（有效期同
    max 规则、source=import），读侧不再"入库成功却看不到"。"""
    from app.db import workspace_models as wm

    db = _db()
    try:
        old_pid, _old_ay_id = _make_history_person(
            db, "2022-2023", "2022-09-01", "2023-07-15",
            "teaching", "秦往", "2022T8-66",
        )
    finally:
        db.close()

    path = _write_grade23_xlsx(
        tmp_path / "F7T接续.xlsx",
        [
            _student_row("2022T8-66", "秦往", class_num=8, chinese=60, math=61,
                         physics=70, physics_grade=71, teaching_label="高二8班(教)"),
        ],
        teaching_col=True,
    )
    prev = _upload(
        client, path, "teaching",
        {
            "teaching_class_id": v1_seed.t8_id,
            "exam_name": EXAM_F7T,
            "exam_date": "2025-11-10",
        },
    )
    assert prev.status_code == 200, prev.text
    item = prev.json()["items"][0]
    assert item["known_students"] == 0
    assert len(item["identity_candidates"]) == 1

    ok = client.post(
        f"{API}/imports/confirm",
        json={
            "token": prev.json()["token"],
            "identity_confirmations": {"2022T8-66": old_pid},
        },
    )
    assert ok.status_code == 200, ok.text
    assert ok.json()["members_synced"] == 1  # 接续成员补入 T8

    db = _db()
    try:
        member = (
            db.query(wm.TeachingClassMember)
            .filter_by(teaching_class_id=v1_seed.t8_id, identity_id=old_pid)
            .one()
        )
        assert member.valid_from == date(2025, 11, 10)  # max 规则，非今天
        assert member.source == "import"
        fact = (
            db.query(wm.ScoreFact)
            .filter_by(exam_name=EXAM_F7T, identity_id=old_pid)
            .one()
        )
        assert fact.score == 70.0
    finally:
        db.close()


# ────────────────────────────── 考试日期冲突（v2.1 边界裁决） ──────────────────────────────


def test_exam_date_conflict_database_and_batch(client, v1_seed, tmp_path):
    """同 exam_name 不同 exam_date：与库内（seed 2025期中=2025-11-06）或
    批内多文件之间 → preview 计 warnings，confirm 整批 409 零写入；
    同名同日期重导不受影响。"""
    from app.db import workspace_models as wm

    # ① 库内冲突：seed 已有 2025期中（2025-11-06），再导 2025-11-20
    path = _write_grade23_xlsx(
        tmp_path / "F日期冲突.xlsx",
        [_student_row("2025H6-81", "秦捌一", chinese=60, math=61)],
    )
    prev = _upload(
        client, path, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": "2025期中", "exam_date": "2025-11-20"},
    )
    assert prev.status_code == 200
    assert any("考试日期冲突" in w for w in prev.json()["items"][0]["warnings"])

    r = client.post(
        f"{API}/imports/confirm", json={"token": prev.json()["token"]}
    )
    assert r.status_code == 409, r.text
    body = r.json()
    assert body["error"] == "link_version_conflict"
    assert body["date_conflicts"][0]["scope"] == "database"
    assert body["date_conflicts"][0]["existing_dates"] == ["2025-11-06"]

    db = _db()
    try:
        assert (
            db.query(wm.ScoreFact)
            .filter(wm.ScoreFact.exam_name == "2025期中")
            .filter(wm.ScoreFact.exam_date == date(2025, 11, 20))
            .count()
            == 0
        )
        batch = db.query(wm.ImportBatch).filter_by(token=prev.json()["token"]).one()
        assert batch.status == "pending"
    finally:
        db.close()

    # ② 同名同日期：与库内一致 → 不冲突，正常入库
    path_same = _write_grade23_xlsx(
        tmp_path / "F日期相同.xlsx",
        [_student_row("2025H6-82", "秦捌二", chinese=62, math=63)],
    )
    prev_same = _upload(
        client, path_same, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": "2025期中", "exam_date": "2025-11-06"},
    )
    assert prev_same.status_code == 200
    ok = client.post(
        f"{API}/imports/confirm", json={"token": prev_same.json()["token"]}
    )
    assert ok.status_code == 200, ok.text
    assert ok.json()["imported"] == 2

    # ③ 批内冲突：两份合法文件文件名推出同考试名、不同日期
    # （"X年Y月期中" → canonical 同名"未知年级未知学期期中考试"，日期取 Y 月 1 日）
    f_nov = _write_grade23_xlsx(
        tmp_path / "2025年11月期中混批A.xlsx",
        [_student_row("2025H6-83", "秦捌三", chinese=64, math=65)],
    )
    f_dec = _write_grade23_xlsx(
        tmp_path / "2025年12月期中混批B.xlsx",
        [_student_row("2025H6-84", "秦捌四", chinese=66, math=67)],
    )
    prev_batch = _upload(client, [f_nov, f_dec], "homeroom", {"class_id": v1_seed.h6_id})
    assert prev_batch.status_code == 200, prev_batch.text
    items = prev_batch.json()["items"]
    assert items[0]["exam_name"] == items[1]["exam_name"]
    assert items[0]["exam_date"] == "2025-11-01"
    assert items[1]["exam_date"] == "2025-12-01"
    for it in items:
        assert any("考试日期冲突" in w and "批内" in w for w in it["warnings"])

    r_batch = client.post(
        f"{API}/imports/confirm", json={"token": prev_batch.json()["token"]}
    )
    assert r_batch.status_code == 409
    body_batch = r_batch.json()
    assert body_batch["date_conflicts"][0]["scope"] == "batch"
    assert body_batch["date_conflicts"][0]["dates"] == ["2025-11-01", "2025-12-01"]

    db = _db()
    try:
        assert (
            db.query(wm.ScoreFact)
            .filter(wm.ScoreFact.exam_name == "未知年级未知学期期中考试")
            .count()
            == 0
        )
    finally:
        db.close()
