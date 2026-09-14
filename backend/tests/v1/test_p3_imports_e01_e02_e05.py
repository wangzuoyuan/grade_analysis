"""P3 导入链路验收：E01 / E02 / E05（契约 p3-imports-analysis.md §1）。

- E01：同一份成绩数据分别 homeroom 导入（全科+总分入库）与 teaching
  导入（仅物理入库）→ /api/v1/scores 两模式各自断言域内字段集合
  （teaching 无 total_type 行、无其他学科）。
- E02：缺考 NULL 行 score=None 不转 0；同 token 重复 confirm → 409；
  重新 preview 同文件 → confirm 幂等（imported=0 skipped=n）。
- E05：confirm 前改成员 → 409 零写入（库内无 ScoreFact 残留）；不同值
  再导入 revise=false → 409 + conflicts；revise=true → 覆写且
  data_revision+1（同 token 重试语义：409 后 batch 保持 pending）。

xlsx 样本在 fixture 内用 openpyxl 程序化生成（高二 3+3 固定列布局，
H/T 解析器共用），不提交二进制。造表 helper 与
test_p3_shared_exams.py 同源复制（conftest 不在本任务允许清单）。
"""

from pathlib import Path

import pytest

API = "/api/v1"
SUBJECT = "物理"

EXAM_H = "P3E1H"
EXAM_T = "P3E1T"
EXAM_E5 = "P3E5H"

# 高二 3+3 固定列布局（列号即 H/T 解析器的 GRADE23_* 常量）：
# 1学号 2班级 3学籍 4姓名 | 5语文 6数学 7英语 | 8物理 9物理等级分 |
# 20 +3总分 | 21主三门总分 22主三门百分位 23学籍排名 | 24 3+3总分 |
# 27语文百分位 28数学百分位 29英语百分位 | 30 教学班标签列（teaching 用）
_HEADERS = {
    1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
    5: "语文", 6: "数学", 7: "英语",
    8: "物理", 9: "物理等级分",
    20: "+3总分", 21: "主三门总分", 22: "主三门百分位", 23: "主三门学籍排名",
    24: "3+3总分", 25: "3+3百分位", 26: "3+3学籍排名",
    27: "语文百分位", 28: "数学百分位", 29: "英语百分位",
}


def _student_row(
    sid, name, *, class_num=6, chinese=None, math=None, english=None,
    physics=None, physics_grade=None, chinese_pct=None,
    plus3=None, main3=None, total3=None, teaching_label=None,
):
    data = {1: sid, 2: class_num, 3: 1, 4: name}
    if chinese is not None:
        data[5] = chinese
    if math is not None:
        data[6] = math
    if english is not None:
        data[7] = english
    if physics is not None:
        data[8] = physics
    if physics_grade is not None:
        data[9] = physics_grade
    if chinese_pct is not None:
        data[27] = chinese_pct
    if plus3 is not None:
        data[20] = plus3
    if main3 is not None:
        data[21] = main3
    if total3 is not None:
        data[24] = total3
    if teaching_label is not None:
        data[30] = teaching_label
    return data


def _write_grade23_xlsx(path: Path, rows, teaching_col: bool = False) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    headers = dict(_HEADERS)
    if teaching_col:
        headers[30] = "教学班"
    for col, text in headers.items():
        ws.cell(row=2, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def _upload(client, path: Path, mode: str, extra: dict | None = None):
    """multipart preview：openpyxl 文件 + 表单作用域/考试覆盖参数。"""
    data = {"mode": mode}
    if extra:
        data.update({k: str(v) for k, v in extra.items()})
    with open(path, "rb") as fh:
        return client.post(
            f"{API}/imports/preview",
            data=data,
            files=[
                (
                    "files",
                    (
                        path.name,
                        fh.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                )
            ],
        )


def _db():
    from app.db.models import SessionLocal

    return SessionLocal()


# 模块内共享：E01 的导入产物供 E02（同模块按定义顺序执行）复用
STATE: dict = {}


@pytest.fixture()
def e1_homeroom_file(tmp_path):
    """E1 同源成绩（homeroom 域学号）：甲乙丙全科 + 总分；丙语文缺考
    （分数空、语文百分位有值 → 行保留 score=None，H 解析器语义）。"""
    return _write_grade23_xlsx(
        tmp_path / "高二第一学期期中-P3E1.xlsx",
        [
            _student_row("2025H6-01", "秦甲", chinese=88, math=92, english=95,
                         physics=90, physics_grade=91, plus3=355, main3=275, total3=436),
            _student_row("2025H6-02", "秦乙", chinese=76, math=81, english=79,
                         physics=84, physics_grade=85),
            _student_row("2025H6-03", "秦丙", chinese=None, chinese_pct=50,
                         math=70, english=72, physics=None),
        ],
    )


@pytest.fixture()
def e1_teaching_file(tmp_path, v1_seed):
    """E1 同源成绩（teaching 域学号，T8 标签列）：戊己为已知成员、
    庚为新学号；丁（T6 标签）为外教学班标签行——应被过滤不入 T8。"""
    from tests.v1.conftest import ALIAS_JI_T, ALIAS_WU_T

    return _write_grade23_xlsx(
        tmp_path / "高二第一学期期中-P3E1T.xlsx",
        [
            _student_row(ALIAS_WU_T, "秦甲", class_num=8, chinese=70, math=71,
                         english=72, physics=78, physics_grade=79,
                         teaching_label="高二8班(教)"),
            _student_row(ALIAS_JI_T, "秦己·T", class_num=8, chinese=66, math=67,
                         english=68, physics=82, physics_grade=83,
                         teaching_label="高二8班(教)"),
            _student_row("2025T8-99", "秦庚", class_num=8, chinese=60, math=61,
                         english=62, physics=60, physics_grade=61,
                         teaching_label="高二8班(教)"),
            _student_row("2025T6-04", "秦丁·T", class_num=6, chinese=55, math=56,
                         english=57, physics=58, physics_grade=59,
                         teaching_label="高二6班(教)"),
        ],
        teaching_col=True,
    )


# ────────────────────────────── E01 ──────────────────────────────


def test_e01_homeroom_import_full_subjects(client, v1_seed, e1_homeroom_file):
    """homeroom 导入：全科 + 总分行入库，grade_score 落列，缺考 NULL。"""
    # 非绑定班 404（契约 §1.1：class_id 显式时校验绑定）
    r404 = _upload(client, e1_homeroom_file, "homeroom", {"class_id": v1_seed.h9_id})
    assert r404.status_code == 404
    assert r404.json()["error"] == "resource_out_of_scope"

    prev = _upload(
        client, e1_homeroom_file, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_H, "exam_date": "2025-11-10"},
    )
    assert prev.status_code == 200, prev.text
    body = prev.json()
    assert body["mode"] == "homeroom"
    assert len(body["items"]) == 1
    item = body["items"][0]
    assert item["kind"] == "student_scores" and item["parsed_ok"]
    assert item["exam_name"] == EXAM_H and item["exam_date"] == "2025-11-10"
    # 甲乙丙均为库内已知（本学年 alias 命中；种子 alias 已带学年标记），
    # 无新学号
    assert item["known_students"] == 3 and item["new_students"] == []
    assert any("缺考" in w for w in item["warnings"])

    ok = client.post(f"{API}/imports/confirm", json={"token": body["token"]})
    assert ok.status_code == 200, ok.text
    result = ok.json()
    # 甲 4 科 + 3 总分；乙 4 科；丙 3 科（物理空无行、无总分）= 14 行
    assert result["imported"] == 14
    assert result["students_created"] == 0
    assert result["exams"] == [{"exam_name": EXAM_H, "exam_date": "2025-11-10"}]
    STATE["e1h_token"] = body["token"]
    STATE["e1h_file"] = e1_homeroom_file

    r = client.get(f"{API}/scores", params={"mode": "homeroom", "exam_name": EXAM_H})
    assert r.status_code == 200
    rows = r.json()["rows"]
    assert len(rows) == 14
    by_key = {(row["person_id"], row["subject"], row["total_type"]): row for row in rows}
    # 全科 + 总分都在；甲物理带等级分（§1.3）；丙语文缺考 NULL 不转 0
    jia_phys = by_key[(v1_seed.jia_h_id, "物理", None)]
    assert jia_phys["score"] == 90.0 and jia_phys["grade_score"] == 91.0
    bing_chinese = by_key[(v1_seed.bing_h_id, "语文", None)]
    assert bing_chinese["score"] is None
    total_types = {row["total_type"] for row in rows if row["total_type"]}
    assert total_types == {"+3", "主三门", "3+3"}
    subjects = {row["subject"] for row in rows if row["subject"]}
    assert subjects == {"语文", "数学", "英语", "物理"}


def test_e01_teaching_import_subject_only(client, v1_seed, e1_teaching_file):
    """teaching 导入：仅物理入库（其他学科列丢弃计 warnings、外标签行
    过滤）；新学号建身份；/scores teaching 无 total_type 键、无其他学科。"""
    prev = _upload(
        client, e1_teaching_file, "teaching",
        {"teaching_class_id": v1_seed.t8_id, "exam_name": EXAM_T, "exam_date": "2025-11-11"},
    )
    assert prev.status_code == 200, prev.text
    item = prev.json()["items"][0]
    assert item["subject"] == SUBJECT
    assert item["class_label"] == "高二8班(教)"
    assert item["known_students"] == 2  # 戊己（历史 alias 命中）
    assert [s["alias"] for s in item["new_students"]] == ["2025T8-99"]
    assert any("已忽略非任教学科列" in w for w in item["warnings"])
    # 外标签行（丁，T6 标签）被过滤，不进本教学班
    assert any("已忽略非本教学班" in w and "2025T6-04" in w for w in item["warnings"])

    ok = client.post(f"{API}/imports/confirm", json={"token": prev.json()["token"]})
    assert ok.status_code == 200, ok.text
    result = ok.json()
    assert result["imported"] == 3  # 戊/己/庚各 1 物理行（丁被标签过滤）
    assert result["students_created"] == 1
    assert result["members_synced"] == 1  # 庚补入 T8；戊己已是成员

    # subject 显式传参与教学班学科一致性：subject 不匹配教学班 → 422
    bad_subject = _upload(
        client, e1_teaching_file, "teaching",
        {"teaching_class_id": v1_seed.t8_id, "subject": "化学", "exam_name": EXAM_T},
    )
    assert bad_subject.status_code == 422
    assert bad_subject.json()["error"] == "invalid_scope_param"

    r = client.get(f"{API}/scores", params={"mode": "teaching", "exam_name": EXAM_T})
    assert r.status_code == 200
    rows = r.json()["rows"]
    assert {row["subject"] for row in rows} == {SUBJECT}  # 恒仅任教学科
    assert all("total_type" not in row for row in rows)  # 键本身缺席
    wu = next(row for row in rows if row["person_id"] == v1_seed.wu_t_id)
    assert wu["score"] == 78.0
    new_scores = [row for row in rows if row["person_id"] not in
                  set(v1_seed.t8_person_ids)]
    assert len(new_scores) == 1 and new_scores[0]["score"] == 60.0

    # S05 前半：T8 无 link，teaching 域考试天然不进 homeroom /scores
    hr = client.get(f"{API}/scores", params={"mode": "homeroom", "exam_name": EXAM_T})
    assert hr.status_code == 200
    assert hr.json()["rows"] == []
    STATE["new_t8_person_id"] = new_scores[0]["person_id"]


# ────────────────────────────── E02 ──────────────────────────────


def test_e02_consumed_token_reimport_idempotent_null_kept(client, v1_seed, tmp_path):
    """同 token 重复 confirm → 409（R4 单次消费）；重新 preview 同文件 →
    confirm 幂等（imported=0 skipped=14）；缺考 NULL 落库不转 0。"""
    # E01 已消费的 token 再次提交 → 409，零副作用
    again = client.post(f"{API}/imports/confirm", json={"token": STATE["e1h_token"]})
    assert again.status_code == 409
    assert again.json()["error"] == "link_version_conflict"

    # 与 E01 同内容重造文件（不依赖前序 tmp 目录存活）
    same = _write_grade23_xlsx(
        tmp_path / "高二第一学期期中-P3E1-again.xlsx",
        [
            _student_row("2025H6-01", "秦甲", chinese=88, math=92, english=95,
                         physics=90, physics_grade=91, plus3=355, main3=275, total3=436),
            _student_row("2025H6-02", "秦乙", chinese=76, math=81, english=79,
                         physics=84, physics_grade=85),
            _student_row("2025H6-03", "秦丙", chinese=None, chinese_pct=50,
                         math=70, english=72, physics=None),
        ],
    )
    prev = _upload(
        client, same, "homeroom",
        {"class_id": v1_seed.h6_id, "exam_name": EXAM_H, "exam_date": "2025-11-10"},
    )
    assert prev.status_code == 200
    # 第二次导入：上一轮新建/接续的学号全部 known
    assert prev.json()["items"][0]["known_students"] == 3

    ok = client.post(f"{API}/imports/confirm", json={"token": prev.json()["token"]})
    assert ok.status_code == 200, ok.text
    result = ok.json()
    assert result["imported"] == 0
    assert result["skipped"] == 14  # 自然键同值幂等
    assert result["revised"] == 0

    # 缺考：库内 score 为 NULL 的行存在（丙语文），绝不转 0
    db = _db()
    try:
        from app.db import workspace_models as wm

        null_rows = (
            db.query(wm.ScoreFact)
            .filter(
                wm.ScoreFact.exam_name == EXAM_H,
                wm.ScoreFact.data_domain == "homeroom",
                wm.ScoreFact.score.is_(None),
            )
            .all()
        )
        assert len(null_rows) == 1
        assert null_rows[0].identity_id == v1_seed.bing_h_id
        assert null_rows[0].subject == "语文"
    finally:
        db.close()


# ────────────────────────────── E05 ──────────────────────────────


@pytest.fixture()
def e5_files(tmp_path):
    """E5 两份文件：v1 正常导入；v2 仅乙物理 84→93（制造值冲突）。"""
    base = [
        _student_row("2025H6-01", "秦甲", chinese=88, math=92, english=95, physics=90, main3=275),
        _student_row("2025H6-02", "秦乙", chinese=76, math=81, english=79, physics=84, main3=236),
        _student_row("2025H6-03", "秦丙", chinese=65, math=70, english=72, physics=68, main3=207),
    ]
    v2 = [dict(row) for row in base]
    v2[1][8] = 93  # 乙物理改分
    return (
        _write_grade23_xlsx(tmp_path / "高二第一学期期中-P3E5.xlsx", base),
        _write_grade23_xlsx(tmp_path / "高二第一学期期中-P3E5-v2.xlsx", v2),
    )


def _yi_enrollment(db, seed):
    from app.db import workspace_models as wm

    return (
        db.query(wm.Enrollment)
        .filter_by(admin_class_id=seed.h6_id, identity_id=seed.yi_h_id)
        .one()
    )


def test_e05_member_drift_atomic_and_revise(client, v1_seed, e5_files):
    from datetime import date, timedelta

    from app.db import workspace_models as wm

    v1_path, v2_path = e5_files
    extra = {"class_id": v1_seed.h6_id, "exam_name": EXAM_E5, "exam_date": "2025-12-01"}

    # ① preview 后成员漂移 → 409 零写入（E05 原子），batch 保持 pending
    prev = _upload(client, v1_path, "homeroom", extra)
    assert prev.status_code == 200
    token = prev.json()["token"]

    db = _db()
    try:
        _yi_enrollment(db, v1_seed).valid_to = date.today() - timedelta(days=1)
        db.commit()
    finally:
        db.close()

    drifted = client.post(f"{API}/imports/confirm", json={"token": token})
    assert drifted.status_code == 409
    assert drifted.json()["error"] == "link_version_conflict"
    assert "member_person_ids" in drifted.json().get("drift", {})

    db = _db()
    try:
        assert db.query(wm.ScoreFact).filter_by(exam_name=EXAM_E5).count() == 0
        batch = db.query(wm.ImportBatch).filter_by(token=token).one()
        assert batch.status == "pending"
        _yi_enrollment(db, v1_seed).valid_to = None  # 恢复成员
        db.commit()
    finally:
        db.close()

    # ② 同 token 重试（恢复后）→ 成功导入 15 行（3 人 × (4 科 + 主三门)）
    ok = client.post(f"{API}/imports/confirm", json={"token": token})
    assert ok.status_code == 200, ok.text
    assert ok.json()["imported"] == 15

    db = _db()
    try:
        yi_phys = (
            db.query(wm.ScoreFact)
            .filter_by(
                exam_name=EXAM_E5, identity_id=v1_seed.yi_h_id, subject="物理"
            )
            .one()
        )
        assert yi_phys.score == 84.0 and yi_phys.data_revision == 1
    finally:
        db.close()

    # ③ 不同值再导入：revise=false → 409 + conflicts 列表，零写入
    prev2 = _upload(client, v2_path, "homeroom", extra)
    assert prev2.status_code == 200
    token2 = prev2.json()["token"]
    conflict = client.post(
        f"{API}/imports/confirm", json={"token": token2, "revise": False}
    )
    assert conflict.status_code == 409
    body = conflict.json()
    assert body["error"] == "link_version_conflict"
    assert body["conflicts"] == [
        {
            "person": "秦乙",
            "subject": "物理",
            "exam_name": EXAM_E5,
            "existing_score": 84.0,
            "new_score": 93.0,
        }
    ]

    db = _db()
    try:
        assert (
            db.query(wm.ScoreFact)
            .filter_by(exam_name=EXAM_E5, identity_id=v1_seed.yi_h_id, subject="物理")
            .one()
            .score
        ) == 84.0  # 拒绝路径零写入
        batch2 = db.query(wm.ImportBatch).filter_by(token=token2).one()
        assert batch2.status == "pending"  # 可重试
    finally:
        db.close()

    # ④ 同 token revise=true 重试 → 覆写 + data_revision+1（E05 修订语义）
    revised = client.post(
        f"{API}/imports/confirm", json={"token": token2, "revise": True}
    )
    assert revised.status_code == 200, revised.text
    result = revised.json()
    assert result["revised"] == 1
    assert result["skipped"] == 14  # 其余同值行幂等
    assert result["imported"] == 0

    db = _db()
    try:
        yi_phys = (
            db.query(wm.ScoreFact)
            .filter_by(
                exam_name=EXAM_E5, identity_id=v1_seed.yi_h_id, subject="物理"
            )
            .one()
        )
        assert yi_phys.score == 93.0
        assert yi_phys.data_revision == 2
    finally:
        db.close()


def test_e05_same_alias_different_name_rejected(client, v1_seed, tmp_path):
    """文件内同学号不同姓名 → preview warning、confirm 整文件 409 拒绝。"""
    bad = _write_grade23_xlsx(
        tmp_path / "撞号.xlsx",
        [
            _student_row("2025H6-01", "秦甲", chinese=88, math=92),
            _student_row("2025H6-01", "秦假", chinese=66, math=70),  # 同学号不同名
        ],
    )
    prev = _upload(client, bad, "homeroom", {"class_id": v1_seed.h6_id, "exam_name": "撞号考"})
    assert prev.status_code == 200
    item = prev.json()["items"][0]
    assert any("同学号不同姓名" in w for w in item["warnings"])

    r = client.post(f"{API}/imports/confirm", json={"token": prev.json()["token"]})
    assert r.status_code == 409
    assert r.json()["error"] == "link_version_conflict"
    assert r.json()["conflicts"]

    db = _db()
    try:
        from app.db import workspace_models as wm

        assert db.query(wm.ScoreFact).filter_by(exam_name="撞号考").count() == 0
    finally:
        db.close()


def test_e05_db_alias_name_conflict_rejected(client, v1_seed, tmp_path):
    """文件学号命中库内 identity 但姓名不同（库内撞号）→ confirm 整文件 409。"""
    bad = _write_grade23_xlsx(
        tmp_path / "库内撞号.xlsx",
        [_student_row("2025H6-01", "秦别人", chinese=88, math=92)],  # 01 号库内是秦甲
    )
    prev = _upload(client, bad, "homeroom", {"class_id": v1_seed.h6_id, "exam_name": "库内撞号考"})
    assert prev.status_code == 200
    assert any("同学号不同姓名" in w for w in prev.json()["items"][0]["warnings"])

    r = client.post(f"{API}/imports/confirm", json={"token": prev.json()["token"]})
    assert r.status_code == 409
    assert r.json()["error"] == "link_version_conflict"

    db = _db()
    try:
        from app.db import workspace_models as wm

        assert db.query(wm.ScoreFact).filter_by(exam_name="库内撞号考").count() == 0
    finally:
        db.close()
