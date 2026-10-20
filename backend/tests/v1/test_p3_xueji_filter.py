"""排名分布类图表的学籍=1 过滤（2026-10-20 用户裁决，班主任工作台全域）。

- 段位分布（bands）与排名分段（rank-distribution）仅总分口径：只统计
  学校成绩明细表学籍=1（闵中学籍）的学生；NULL/3/4 均不参与。
- 排名频次（rank-frequency）与排名区间（rank-range）：仅 total_rank
  口径过滤；单科（百分位/等级分）口径保留全部学生。
- 学籍存身份上、查询时按当前值过滤——历史考试图表随学籍更新即时联动。
- 导入即更新：明细表文件学籍列有值即覆盖身份 xueji（不再仅补空）。
"""

from datetime import date
from pathlib import Path
from urllib.parse import quote

import pytest

API = "/api/v1"
EXAM = "P3XJ学籍过滤考"
EXAM_NO_RANK = "P3XJ无名次考"

_HEADERS = {
    1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
    5: "语文", 6: "数学", 7: "英语",
    20: "+3总分", 21: "主三门总分", 23: "主三门学籍排名",
}


def _write_detail(path: Path, rows) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col, text in _HEADERS.items():
        ws.cell(row=2, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def _upload(client, path: Path, extra: dict):
    return client.post(
        f"{API}/imports/preview",
        data={"mode": "homeroom", **{k: str(v) for k, v in extra.items()}},
        files=[("files", (path.name, path.read_bytes(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
    )


@pytest.fixture(scope="module")
def xueji_seed(v1_seed):
    """甲=1（应计入）、乙=3（不计入）、丙=4（不计入）；主三门带名次。"""
    from app.db import workspace_models as wm
    from app.db.models import SessionLocal

    s = v1_seed
    db = SessionLocal()

    def set_xueji(ident_id, value):
        ident = db.get(wm.WsStudentIdentity, ident_id)
        ident.xueji = value

    set_xueji(s.jia_h_id, 1)
    set_xueji(s.yi_h_id, 3)
    set_xueji(s.bing_h_id, 4)

    def fact(ident_id, exam, total_type, score, rank, subject=None, percentile=None):
        f = wm.ScoreFact()
        f.data_domain = "homeroom"
        f.academic_year_id = s.ay_id
        f.exam_name = exam
        f.exam_date = date(2025, 12, 10)
        f.class_ref_id = s.h6_id
        f.identity_id = ident_id
        f.subject = subject
        f.total_type = total_type
        f.score = score
        f.xueji_rank = rank
        f.grade_percentile = percentile
        f.source = "p3-xueji-test"
        db.add(f)

    # 主三门总分行带名次：甲 50 名、乙 300 名、丙 450 名
    fact(s.jia_h_id, EXAM, "主三门", 275.0, 50)
    fact(s.yi_h_id, EXAM, "主三门", 236.0, 300)
    fact(s.bing_h_id, EXAM, "主三门", 207.0, 450)
    # 语文单科行：带百分位（频次分箱用）与名次（区间筛选按名次匹配用）
    fact(s.jia_h_id, EXAM, None, 88.0, 10, subject="语文", percentile=0.9)
    fact(s.yi_h_id, EXAM, None, 76.0, 20, subject="语文", percentile=0.5)
    fact(s.bing_h_id, EXAM, None, 65.0, 30, subject="语文", percentile=0.3)
    # 无名次考试（bands 应 409）
    fact(s.jia_h_id, EXAM_NO_RANK, "主三门", 275.0, None)

    db.commit()
    yield
    db.close()


def _q(client, path, **params):
    r = client.get(path, params=params)
    return r


def test_bands_and_distribution_count_only_xueji1(client, xueji_seed, v1_seed):
    s = v1_seed
    # 段位分布：阈值默认 高分1-80/临界400-500/薄弱501+ → 只有甲(50名)落高分段
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name=EXAM, subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_type"] == "主三门"
    # bands 顺序固定：高分段/临界段/薄弱段
    labels = [b["label"].split("（")[0] for b in body["bands"]]
    assert labels == ["高分段", "临界段", "薄弱段"]
    counts = {b["label"].split("（")[0]: b["count"] for b in body["bands"]}
    assert counts["高分段"] == 1  # 仅甲（乙=3、丙=4 不计）
    assert sum(counts.values()) == 1

    # 排名分段（顶部柱状图）：40 名一档，只有甲的 41-80 档计 1
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-distribution",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    dist = r.json()
    main3 = next(se for se in dist["series"] if se["total_type"] == "主三门")
    assert sum(main3["counts"].values()) == 1
    assert main3["counts"].get("r41_80") == 1
    assert main3["counts"].get("r281_320", 0) == 0  # 乙 300 名不计
    assert main3["counts"].get("r441_480", 0) == 0  # 丙 450 名不计


def test_frequency_filters_total_but_not_subject(client, xueji_seed, v1_seed):
    s = v1_seed
    # 总分名次口径：仅甲
    r = _q(client, f"{API}/homeroom/analysis/rank-frequency",
           metric="total:主三门", exam_names=EXAM,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    body = r.json()
    students = body["students"]
    assert [st["name"] for st in students] == ["秦甲"]

    # 单科百分位口径：全部学生（不过滤）
    r = _q(client, f"{API}/homeroom/analysis/rank-frequency",
           metric="subject:语文", exam_names=EXAM,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    names = sorted(st["name"] for st in r.json()["students"])
    assert names == ["秦丙", "秦乙", "秦甲"]


def test_rank_range_filters_total_but_not_subject(client, xueji_seed, v1_seed):
    s = v1_seed
    # 总分口径 1-500 名：仅甲乙丙中学籍=1 的甲
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-range",
           metric="total:主三门", rank_min=1, rank_max=500,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    assert [st["name"] for st in r.json()["students"]] == ["秦甲"]

    # 单科口径 1-100 名区间：全部学生（不过滤）
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-range",
           metric="subject:语文", rank_min=1, rank_max=100,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    names = sorted(st["name"] for st in r.json()["students"])
    assert names == ["秦丙", "秦乙", "秦甲"]


def test_bands_409_when_no_rank_or_no_xueji1(client, xueji_seed, v1_seed):
    """无名次考试 → 409（F10 语义保留）；人口全非学籍=1 同样 409。"""
    s = v1_seed
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name=EXAM_NO_RANK, subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 409

    # 甲改为学籍 4 后：唯一学籍=1 的人退出 → 历史考试图表即时联动为 409
    from app.db import workspace_models as wm
    from app.db.models import SessionLocal

    db = SessionLocal()
    db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 4
    db.commit()
    db.close()
    try:
        r = _q(client, f"{API}/homeroom/analysis/bands",
               exam_name=EXAM, subject="主三门", metric="total",
               academic_year_id=s.ay_id, class_id=s.h6_id)
        assert r.status_code == 409
        r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-distribution",
               academic_year_id=s.ay_id, class_id=s.h6_id)
        assert r.status_code == 200
        main3 = next(se for se in r.json()["series"] if se["total_type"] == "主三门")
        assert sum(main3["counts"].values()) == 0
    finally:
        db = SessionLocal()
        db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 1
        db.commit()
        db.close()


def test_import_updates_xueji_from_file(client, v1_seed, tmp_path):
    """明细表导入：文件学籍列有值即覆盖身份 xueji（既有值也会被刷新）。"""
    from app.db import workspace_models as wm
    from app.db.models import SessionLocal
    from tests.v1.conftest import ALIAS_JIA_H

    s = v1_seed
    db = SessionLocal()
    db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 1
    db.commit()
    db.close()

    detail = _write_detail(tmp_path / "学籍更新明细.xlsx", [
        {1: ALIAS_JIA_H, 2: 6, 3: 3, 4: "秦甲", 5: 90, 20: 176, 21: 276, 23: 40},  # 文件学籍=3
    ])
    preview = _upload(client, detail, {
        "class_id": s.h6_id, "exam_name": "P3XJ学籍更新考", "exam_date": "2025-12-11",
    })
    assert preview.status_code == 200, preview.text
    confirmed = client.post(f"{API}/imports/confirm", json={"token": preview.json()["token"]})
    assert confirmed.status_code == 200, confirmed.text

    db = SessionLocal()
    assert db.get(wm.WsStudentIdentity, s.jia_h_id).xueji == 3  # 覆盖为文件值
    db.close()
    # 学籍=1 人口随之联动：新考试 bands 409（甲不再计入）
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name="P3XJ学籍更新考", subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 409

    db = SessionLocal()
    db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 1
    db.commit()
    db.close()
