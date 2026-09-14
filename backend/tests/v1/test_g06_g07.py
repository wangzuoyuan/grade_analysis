""" 二轮审核 G06/G07 回归（契约 p3-imports-analysis.md §1.1 v2.2 / §1.4.1 v2.2）。

- G06（NULL 学年 alias 不再自动混档 + 接续必补学籍）：
  * 未标学年（academic_year_id NULL）且无有效期证据的登记 → 导入同号
    同名只出 identity_candidates（basis='unmarked_alias'），绝不自动
    接续；confirm 无确认 → 409 零写入；显式确认 → 接续旧 identity 并
    补建目标行政班 Enrollment（F07，任何接续路径都补学籍）。
  * NULL 学年但 valid_from/valid_to 明确覆盖目标文件日期 → 直接 known
    （可证实该号事发时仍在册）；confirm 后同样补学籍。
  * NULL 学年但有效期早于文件日期（已离册）→ 不命中，进候选。
  * 跨届同号不同名（未标学年命中）→ 撞号整文件 409 拒绝。
  * 确认伪造 person（未持有该 alias 的任何历史/未标学年登记）→ 422。
- G07（历史时点不看当前 status）：
  * 经正式 POST archive（transferred/graduated, valid_to=2026-08-01）
    离班后，2025-11-06 历史考试的 stats/students 人群不变（含该生）；
    /homeroom/students 当前名册不再含该生。
  * 共享投影（有 link 时）：H-only 历史考试从 teaching 统计/考试清单
    仍可读（考试时点成员交集不看 status）。
  * 考后入班（valid_from 晚于考试日）不进历史考试人群。

在 tests/v1/conftest.py 合成样本（v1_seed）之上 ORM 直种场景；所有
成员状态/新增行改动均在 finally 恢复或清理，不污染同模块其他用例。
造表 helper 与 test_p3_imports_e01_e02_e05.py 同源复制。
"""

from datetime import date
from pathlib import Path
from urllib.parse import quote

import pytest

API = "/api/v1"

EXAM_G6 = "G6接续考"
EXAM_G6_KNOWN = "G6在册考"
EXAM_G6_EXPIRED = "G6过档考"
EXAM_G6_COLLIDE = "G6撞号考"
EXAM_E1 = "2025期中"  # v1_seed 原样：exam_date=2025-11-06
EXAM_G7_PROJ = "G7投影考"

_HEADERS = {
    1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
    5: "语文", 6: "数学", 7: "英语",
    8: "物理", 9: "物理等级分",
    20: "+3总分", 21: "主三门总分", 22: "主三门百分位", 23: "主三门学籍排名",
    24: "3+3总分", 25: "3+3百分位", 26: "3+3学籍排名",
    27: "语文百分位", 28: "数学百分位", 29: "英语百分位",
}


def _student_row(sid, name, *, class_num=6, chinese=None, math=None):
    data = {1: sid, 2: class_num, 3: 1, 4: name}
    if chinese is not None:
        data[5] = chinese
    if math is not None:
        data[6] = math
    return data


def _write_grade23_xlsx(path: Path, rows) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col, text in _HEADERS.items():
        ws.cell(row=2, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def _upload(client, path: Path, extra: dict | None = None):
    data = {"mode": "homeroom"}
    if extra:
        data.update({k: str(v) for k, v in extra.items()})
    with open(path, "rb") as fh:
        return client.post(
            f"{API}/imports/preview",
            data=data,
            files=[
                (
                    "files",
                    (
                        path.name,
                        fh.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                )
            ],
        )


def _db():
    from app.db.models import SessionLocal

    return SessionLocal()


def _make_unmarked_person(db, display_name: str, alias_value: str, *,
                          valid_from=None, valid_to=None):
    """造一个"未标学年（NULL）"的 H 域历史登记（迁移数据的常见形态）。

    valid_from/valid_to 可选：None 即无任何有效期证据（v2.2/G06 下只能
    进候选）；给显式日期则按自身有效期参与"覆盖文件日期"判定。
    返回 identity_id。"""
    from app.db import workspace_models as wm

    ident = wm.WsStudentIdentity(data_domain="homeroom", display_name=display_name)
    db.add(ident)
    db.flush()
    db.add(
        wm.WsStudentAlias(
            identity_id=ident.id,
            alias_value=alias_value,
            data_domain="homeroom",
            academic_year_id=None,
            alias_scope="none",  # 冗余列与 NULL 学年同步（唯一约束）
            source="g06-test",
            valid_from=valid_from,
            valid_to=valid_to,
        )
    )
    db.commit()
    return ident.id


def _purge_created(db, exam_names, person_ids):
    """G06 场景收尾：删除本用例写入的考试事实/学籍/别名/身份，恢复模块
    共享 seed 状态。G07 用例依赖 E1 的考试日期解析（本班全场景最大
    exam_date）与 3 人 cohort，G6 遗留行会改写两者。"""
    from app.db import workspace_models as wm

    if exam_names:
        db.query(wm.ScoreFact).filter(
            wm.ScoreFact.exam_name.in_(list(exam_names))
        ).delete(synchronize_session=False)
    for pid in person_ids:
        db.query(wm.Enrollment).filter_by(identity_id=pid).delete(
            synchronize_session=False
        )
        db.query(wm.WsStudentAlias).filter_by(identity_id=pid).delete(
            synchronize_session=False
        )
        db.query(wm.WsStudentIdentity).filter_by(id=pid).delete(
            synchronize_session=False
        )
    db.commit()


# ────────────────────────────── G06：NULL 学年 alias ──────────────────────────────


def test_g06_unmarked_alias_candidate_then_confirm_backfills_enrollment(
    client, v1_seed, tmp_path
):
    """未标学年且无有效期证据：preview 出候选（basis=unmarked_alias）→
    无确认 409 零写入 → 伪造确认 422 → 显式确认接续旧 identity + 补建
    目标行政班 Enrollment（任何接续路径都补学籍，G06）。"""
    from app.db import workspace_models as wm

    db = _db()
    try:
        old_pid = _make_unmarked_person(db, "秦旧", "2024H6-88")
    finally:
        db.close()

    try:
        path = _write_grade23_xlsx(
            tmp_path / "G6接续.xlsx",
            [_student_row("2024H6-88", "秦旧", chinese=77, math=88)],
        )
        prev = _upload(
            client, path,
            {"class_id": v1_seed.h6_id, "exam_name": EXAM_G6, "exam_date": "2025-11-10"},
        )
        assert prev.status_code == 200, prev.text
        item = prev.json()["items"][0]
        # 未标学年不算已知身份：不算 known、不算 new，只出候选
        assert item["known_students"] == 0
        assert item["new_students"] == []
        assert item["identity_candidates"] == [
            {
                "alias": "2024H6-88",
                "name": "秦旧",
                "person_id": old_pid,
                "academic_year_id": None,
                "academic_year_name": None,
                "basis": "unmarked_alias",
            }
        ]
        assert any("未标学年" in w or "历史学年" in w for w in item["warnings"])

        # 未逐项确认 → 409 + 完整候选清单，零写入，batch 保持 pending
        denied = client.post(
            f"{API}/imports/confirm", json={"token": prev.json()["token"]}
        )
        assert denied.status_code == 409, denied.text
        body = denied.json()
        assert body["error"] == "link_version_conflict"
        assert body["identity_candidates"] == item["identity_candidates"]

        db = _db()
        try:
            assert db.query(wm.ScoreFact).filter_by(exam_name=EXAM_G6).count() == 0
            assert (
                db.query(wm.Enrollment)
                .filter_by(admin_class_id=v1_seed.h6_id, identity_id=old_pid)
                .count()
                == 0
            )
            batch = db.query(wm.ImportBatch).filter_by(token=prev.json()["token"]).one()
            assert batch.status == "pending"
        finally:
            db.close()

        # 确认到从未持有该号的人（乙）→ 422；未标学年登记也是合法持有证据，
        # 但乙压根没有任何形式的登记
        forged = client.post(
            f"{API}/imports/confirm",
            json={
                "token": prev.json()["token"],
                "identity_confirmations": {"2024H6-88": v1_seed.yi_h_id},
            },
        )
        assert forged.status_code == 422, forged.text

        # 同 token 显式确认 → 接续旧 identity，F07 补学籍
        ok = client.post(
            f"{API}/imports/confirm",
            json={
                "token": prev.json()["token"],
                "identity_confirmations": {"2024H6-88": old_pid},
            },
        )
        assert ok.status_code == 200, ok.text
        result = ok.json()
        assert result["imported"] == 2
        assert result["students_created"] == 0  # 接续既有人，不是新建
        assert result["members_synced"] == 1  # F07：补建目标行政班学籍

        db = _db()
        try:
            fact = (
                db.query(wm.ScoreFact).filter_by(exam_name=EXAM_G6, subject="语文").one()
            )
            assert fact.identity_id == old_pid  # 成绩落到历史 identity
            enrollment = (
                db.query(wm.Enrollment)
                .filter_by(admin_class_id=v1_seed.h6_id, identity_id=old_pid)
                .one()
            )
            # F07 有效期 = max(学年 start 2025-09-01, 文件日期 2025-11-10)
            assert enrollment.valid_from == date(2025, 11, 10)
            assert enrollment.status == "active"
            # 本学年 alias 已补登
            assert (
                db.query(wm.WsStudentAlias)
                .filter_by(
                    alias_value="2024H6-88",
                    data_domain="homeroom",
                    academic_year_id=v1_seed.ay_id,
                )
                .one()
                .identity_id
                == old_pid
            )
        finally:
            db.close()

        # 读侧名册可见："导入成功却看不到"已消除
        roster = client.get(
            f"{API}/homeroom/students",
            params={"academic_year_id": v1_seed.ay_id, "class_id": v1_seed.h6_id},
        )
        assert roster.status_code == 200, roster.text
        assert old_pid in {s["person_id"] for s in roster.json()["students"]}
    finally:
        db = _db()
        try:
            _purge_created(db, [EXAM_G6], [old_pid])
        finally:
            db.close()


def test_g06_unmarked_alias_with_validity_covering_file_date_is_known(
    client, v1_seed, tmp_path
):
    """NULL 学年但 valid_from/valid_to 明确覆盖目标文件日期（可证实事发
    时仍在册）→ 直接 known；confirm 后同样补建目标班学籍（合法接续路径
    一律验证目标班关系）。"""
    from app.db import workspace_models as wm

    db = _db()
    try:
        pid = _make_unmarked_person(
            db, "秦在册", "2024H6-77",
            valid_from=date(2024, 9, 1), valid_to=date(2026, 7, 15),
        )
    finally:
        db.close()

    try:
        path = _write_grade23_xlsx(
            tmp_path / "G6在册.xlsx",
            [_student_row("2024H6-77", "秦在册", chinese=71, math=72)],
        )
        prev = _upload(
            client, path,
            {"class_id": v1_seed.h6_id, "exam_name": EXAM_G6_KNOWN,
             "exam_date": "2025-11-10"},
        )
        assert prev.status_code == 200, prev.text
        item = prev.json()["items"][0]
        assert item["known_students"] == 1  # 有效期覆盖文件日期 → 已知身份
        assert item["identity_candidates"] == []
        assert item["new_students"] == []

        ok = client.post(
            f"{API}/imports/confirm", json={"token": prev.json()["token"]}
        )
        assert ok.status_code == 200, ok.text
        assert ok.json()["students_created"] == 0
        assert ok.json()["members_synced"] == 1  # 已知接续同样补学籍

        db = _db()
        try:
            fact = (
                db.query(wm.ScoreFact)
                .filter_by(exam_name=EXAM_G6_KNOWN, subject="语文")
                .one()
            )
            assert fact.identity_id == pid
            enrollment = (
                db.query(wm.Enrollment)
                .filter_by(admin_class_id=v1_seed.h6_id, identity_id=pid)
                .one()
            )
            assert enrollment.valid_from == date(2025, 11, 10)
        finally:
            db.close()
    finally:
        db = _db()
        try:
            _purge_created(db, [EXAM_G6_KNOWN], [pid])
        finally:
            db.close()


def test_g06_unmarked_alias_expired_or_undated_goes_candidate(
    client, v1_seed, tmp_path
):
    """NULL 学年但有效期早于文件日期（已离册）/完全无日期 → 均不命中
    known，一律并入候选（v2.2：未知学年不是已证明的本学年身份）。"""
    db = _db()
    try:
        expired_pid = _make_unmarked_person(
            db, "秦过档", "2024H6-65",
            valid_from=date(2024, 9, 1), valid_to=date(2025, 8, 31),
        )
        undated_pid = _make_unmarked_person(db, "秦无档", "2024H6-64")
    finally:
        db.close()

    try:
        path = _write_grade23_xlsx(
            tmp_path / "G6过档.xlsx",
            [
                _student_row("2024H6-65", "秦过档", chinese=65, math=66),
                _student_row("2024H6-64", "秦无档", chinese=64, math=63),
            ],
        )
        prev = _upload(
            client, path,
            {"class_id": v1_seed.h6_id, "exam_name": EXAM_G6_EXPIRED,
             "exam_date": "2025-11-10"},
        )
        assert prev.status_code == 200, prev.text
        item = prev.json()["items"][0]
        assert item["known_students"] == 0
        assert item["new_students"] == []
        got = {
            (c["alias"], c["person_id"], c["basis"])
            for c in item["identity_candidates"]
        }
        assert got == {
            ("2024H6-65", expired_pid, "unmarked_alias"),
            ("2024H6-64", undated_pid, "unmarked_alias"),
        }

        # 两个候选都未确认 → 409 零写入
        denied = client.post(
            f"{API}/imports/confirm", json={"token": prev.json()["token"]}
        )
        assert denied.status_code == 409
        assert len(denied.json()["identity_candidates"]) == 2
    finally:
        db = _db()
        try:
            _purge_created(db, [EXAM_G6_EXPIRED], [expired_pid, undated_pid])
        finally:
            db.close()


def test_g06_unmarked_alias_cross_cohort_same_number_different_name_rejected(
    client, v1_seed, tmp_path
):
    """未标学年命中 + 文件姓名不同（跨届回收号）→ 撞号防呆：preview 计
    warning，confirm 整文件 409 拒绝（优先于候选确认），不自动改名合并。"""
    db = _db()
    try:
        owner_pid = _make_unmarked_person(db, "秦老号", "2024H6-66")
    finally:
        db.close()

    try:
        path = _write_grade23_xlsx(
            tmp_path / "G6撞号.xlsx",
            [_student_row("2024H6-66", "秦别人", chinese=66, math=77)],
        )
        prev = _upload(
            client, path,
            {"class_id": v1_seed.h6_id, "exam_name": EXAM_G6_COLLIDE,
             "exam_date": "2025-11-12"},
        )
        assert prev.status_code == 200
        item = prev.json()["items"][0]
        assert any("同学号不同姓名" in w for w in item["warnings"])

        r = client.post(
            f"{API}/imports/confirm", json={"token": prev.json()["token"]}
        )
        assert r.status_code == 409
        body = r.json()
        assert body["error"] == "link_version_conflict"
        assert body["conflicts"]  # 撞号冲突，而非候选清单
        assert "identity_candidates" not in body
    finally:
        db = _db()
        try:
            _purge_created(db, [EXAM_G6_COLLIDE], [owner_pid])
        finally:
            db.close()


# ────────────────────────────── G07：历史时点不看当前 status ──────────────────────────────


def _archive(client, seed, person_id: int, status: str, valid_to: str):
    return client.post(
        f"{API}/homeroom/students/{person_id}/archive",
        json={"status": status, "valid_to": valid_to},
        params={"academic_year_id": seed.ay_id, "class_id": seed.h6_id},
    )


def _restore_enrollment(db, seed, person_id: int):
    """archive 场景收尾：恢复 seed 学籍行原状（active、无离班日）。"""
    from app.db import workspace_models as wm

    enr = (
        db.query(wm.Enrollment)
        .filter_by(admin_class_id=seed.h6_id, identity_id=person_id)
        .one()
    )
    enr.status = "active"
    enr.valid_to = None
    db.commit()


def _history_exam(client, seed, exam_name: str):
    return client.get(
        f"{API}/homeroom/analysis/exams/{quote(exam_name)}/stats",
        params={"academic_year_id": seed.ay_id, "class_id": seed.h6_id},
    )


def _history_students(client, seed, exam_name: str):
    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote(exam_name)}/students",
        params={"academic_year_id": seed.ay_id, "class_id": seed.h6_id},
    )
    assert r.status_code == 200, r.text
    return {s["person_id"] for s in r.json()["students"]}


def _current_roster(client, seed):
    r = client.get(
        f"{API}/homeroom/students",
        params={"academic_year_id": seed.ay_id, "class_id": seed.h6_id},
    )
    assert r.status_code == 200, r.text
    return r.json()["students"]


def test_g07_archived_transferred_keeps_history_exam_out_of_current_roster(
    client, v1_seed
):
    """正式 archive 接口离班（transferred, valid_to=2026-08-01）后：
    2025-11-06 历史考试的人群/均分不变（甲仍在 cohort）；当前名册不含
    甲。绝不通过让离班记录伪装 active 修历史。"""
    r = _archive(client, v1_seed, v1_seed.jia_h_id, "transferred", "2026-08-01")
    assert r.status_code == 200, r.text
    assert r.json() == {
        "person_id": v1_seed.jia_h_id,
        "status": "transferred",
        "valid_to": "2026-08-01",
    }

    try:
        body = _history_exam(client, v1_seed, EXAM_E1).json()
        assert body["metadata"]["membership_basis"] == "exam"
        assert body["cohort_size"] == 3  # 甲实际参加了该场考试
        chinese = next(x for x in body["subjects"] if x["subject"] == "语文")
        assert chinese["avg"] == 76.33 and chinese["valid_count"] == 3
        assert v1_seed.jia_h_id in _history_students(client, v1_seed, EXAM_E1)

        roster_ids = {s["person_id"] for s in _current_roster(client, v1_seed)}
        assert v1_seed.jia_h_id not in roster_ids  # 当前名册：已离班
        assert v1_seed.yi_h_id in roster_ids and v1_seed.bing_h_id in roster_ids
    finally:
        db = _db()
        try:
            _restore_enrollment(db, v1_seed, v1_seed.jia_h_id)
        finally:
            db.close()

    # 恢复后当前名册复位（不污染同模块其他用例）
    assert v1_seed.jia_h_id in {
        s["person_id"] for s in _current_roster(client, v1_seed)
    }


def test_g07_archived_graduated_same_semantics(client, v1_seed):
    """graduated 同理：历史考试人群不变，当前名册剔除。"""
    r = _archive(client, v1_seed, v1_seed.yi_h_id, "graduated", "2026-08-01")
    assert r.status_code == 200, r.text

    try:
        body = _history_exam(client, v1_seed, EXAM_E1).json()
        assert body["cohort_size"] == 3
        chinese = next(x for x in body["subjects"] if x["subject"] == "语文")
        assert chinese["valid_count"] == 3
        assert v1_seed.yi_h_id in _history_students(client, v1_seed, EXAM_E1)

        assert v1_seed.yi_h_id not in {
            s["person_id"] for s in _current_roster(client, v1_seed)
        }
    finally:
        db = _db()
        try:
            _restore_enrollment(db, v1_seed, v1_seed.yi_h_id)
        finally:
            db.close()


def test_g07_joined_after_exam_not_in_history(client, v1_seed):
    """考后入班（valid_from=2025-11-10 晚于考试 11-06）：历史考试人群
    不扩；当前名册含该生。"""
    from app.db import workspace_models as wm

    db = _db()
    late = wm.WsStudentIdentity(data_domain="homeroom", display_name="秦后")
    db.add(late)
    db.flush()
    db.add(
        wm.Enrollment(
            admin_class_id=v1_seed.h6_id,
            identity_id=late.id,
            status="active",
            valid_from=date(2025, 11, 10),
        )
    )
    db.commit()
    try:
        body = _history_exam(client, v1_seed, EXAM_E1).json()
        assert body["cohort_size"] == 3  # 仍是甲乙丙
        assert _history_students(client, v1_seed, EXAM_E1) == set(
            v1_seed.h_person_ids
        )
        assert late.id in {
            s["person_id"] for s in _current_roster(client, v1_seed)
        }
    finally:
        db.query(wm.Enrollment).filter_by(identity_id=late.id).delete(
            synchronize_session=False
        )
        db.query(wm.WsStudentIdentity).filter_by(id=late.id).delete(
            synchronize_session=False
        )
        db.commit()
        db.close()


def test_g07_shared_projection_keeps_archived_student_for_history_exam(
    client, v1_seed
):
    """共享投影（有 link 时）：甲经正式 archive 离班后，其 H-only 历史
    考试事实仍从 teaching 可读（考试时点成员交集只看有效期，不看当前
    status）；考试清单同步保留。"""
    from app.db import workspace_models as wm

    db = _db()
    db.add(
        wm.ScoreFact(
            data_domain="homeroom",
            academic_year_id=v1_seed.ay_id,
            exam_name=EXAM_G7_PROJ,
            exam_date=date(2026, 1, 10),
            class_ref_id=v1_seed.h6_id,
            identity_id=v1_seed.jia_h_id,
            subject="物理",
            score=66.0,
            source="g07-test",
        )
    )
    db.commit()
    db.close()

    def _teaching_stats():
        r = client.get(
            f"{API}/teaching/analysis/exams/{quote(EXAM_G7_PROJ)}/stats",
            params={
                "academic_year_id": v1_seed.ay_id,
                "teaching_class_id": v1_seed.t6_id,
            },
        )
        assert r.status_code == 200, r.text
        return r.json()

    # 离班前：反向投影可读（基线）
    base = _teaching_stats()
    assert base["valid_count"] == 1 and base["avg"] == 66.0
    assert base["cohort_size"] == 3

    r = _archive(client, v1_seed, v1_seed.jia_h_id, "transferred", "2026-08-01")
    assert r.status_code == 200, r.text

    try:
        # 离班后：历史考试投影不变（2026-01-10 时甲双侧均在册）
        after = _teaching_stats()
        assert after["valid_count"] == 1 and after["avg"] == 66.0
        assert after["cohort_size"] == 3

        # 考试清单（统一可读口径）仍列出该考试
        listed = client.get(
            f"{API}/shared/exams",
            params={
                "mode": "teaching",
                "academic_year_id": v1_seed.ay_id,
                "teaching_class_id": v1_seed.t6_id,
            },
        )
        assert listed.status_code == 200, listed.text
        assert EXAM_G7_PROJ in {e["exam_name"] for e in listed.json()["exams"]}
    finally:
        db = _db()
        try:
            _restore_enrollment(db, v1_seed, v1_seed.jia_h_id)
            db.query(wm.ScoreFact).filter_by(
                exam_name=EXAM_G7_PROJ, source="g07-test"
            ).delete(synchronize_session=False)
            db.commit()
        finally:
            db.close()
