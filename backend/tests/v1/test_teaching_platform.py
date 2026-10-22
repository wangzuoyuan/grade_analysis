"""教学平台纵向回归：只使用合成 Excel 和隔离数据库。"""
from datetime import date
from io import BytesIO
from types import SimpleNamespace
import itertools

import pytest
from openpyxl import Workbook, load_workbook
from app.db.models import SessionLocal, Teacher, engine
from app.db.workspace_models import (AcademicYear,TeachingClass,TeachingClassMember,ScoreFact,
    TeachingExamScope,TeachingClassQuality,TeachingComparisonRoster,WsStudentIdentity,WsStudentAlias,ImportBatch)
from app.ingest.teaching_platform import GRADE_LABELS,parse

API="/api/v1/teaching/platform"
counter=itertools.count(1)


def xlsx(rows,merges=()):
    w=Workbook()
    for row in rows:
        w.active.append(row)
    for area in merges:
        w.active.merge_cells(area)
    b=BytesIO();w.save(b)
    return b.getvalue()


def detail(grade=2,rows=None):
    return xlsx([['学号','姓名',*(['分数','学籍排名','年级百分位'] if grade==1 else ['原始分数','等级分数'])],
        *(rows if rows is not None else [['SYN001','合成甲',60,None,None] if grade==1 else ['SYN001','合成甲',60,None]])])


def quality(grade=2,kind="teaching",labels=None,teachers=None,avg=50):
    labels=labels or ['合成A','合成a']
    teachers=teachers if teachers is not None else ['合成教师']+[None]*(len(labels)-1)
    prefix=['班级类型'] if kind=='administrative' else []
    top=prefix+['班级','任课教师' if prefix else '任教老师','人数',None]
    sub=prefix+['班级','任课教师' if prefix else '任教老师','应考','实考']
    if grade==1:
        top+=['总分','平均分','标准差','最高分','最低分']
        sub+=['总分','平均分','标准差','最高分','最低分']
        bands=['优秀人数（100-85）','良好人数（84-75）','合格人数（74-60）','低分人数（40分以下）']
    else:
        top+=['平均分',None,'标准差']
        sub+=['原始','等级','标准差']
        bands=list(GRADE_LABELS.values())
    for b in bands:
        top += [b if grade==1 else '等级人数及比率',None]
        sub += ['人数' if grade==1 else b,'比率']
    rows=[top,sub]
    for label,teacher in zip(labels,teachers):
        row=(['合成班型'] if prefix else [])+[label,teacher,10,10]
        row += [avg*10,avg,8,90,10] if grade==1 else [avg,55,8]
        for b in bands:
            row += [10 if b=='C+' else 0,1 if b=='C+' else 0]
        rows.append(row)
    from openpyxl.utils import get_column_letter as col
    merges=[f'{col(len(prefix)+3)}1:{col(len(prefix)+4)}1']
    if grade!=1:
        merges.append(f'{col(len(prefix)+5)}1:{col(len(prefix)+6)}1')
    return xlsx(rows,merges)


@pytest.fixture
def env(client):
    from app.db.models import Base
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    n=next(counter)
    db=SessionLocal()
    db.merge(Teacher(id=1,name='合成教师'))
    ay=AcademicYear(name=f'合成学年{n}',start_date=date(2030,9,1),end_date=date(2031,8,31))
    db.add(ay);db.flush()
    tc=TeachingClass(academic_year_id=ay.id,subject='物理',label='合成A')
    db.add(tc);db.commit()
    result=SimpleNamespace(client=client,db=db,year=ay.id,tc=tc.id)
    yield result
    db.close()


def preview(e,files=None,grade=2,operation='initial',target=True,name='合成期中',**extra):
    form={'academic_year_id':str(e.year),'grade':str(grade),'subject':'物理','exam_name':name,'exam_date':'2030-11-01','operation':operation,**extra}
    if target:form['teaching_class_id']=str(e.tc)
    r=e.client.post(API+'/imports/preview',data=form,files=[('files',(n,b,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')) for n,b in (files or [('明细.xlsx',detail(grade))])])
    assert r.status_code==200,r.text
    return r.json()


def confirm(e,p,status=200,**extra):
    r=e.client.post(API+'/imports/confirm',json={'token':p['token'],'accept_roster_changes':True,**extra})
    assert r.status_code==status,r.text
    e.db.expire_all()
    return r.json()


def test_five_formats_and_selection():
    assert parse(detail(1),'高一.xlsx',1,'物理')['rows'][0]['score']==60
    assert parse(detail(),'高二.xlsx',2,'物理')['rows'][0]['grade_score'] is None
    high1=parse(quality(1,'administrative',labels=['01','02'],teachers=['合成教师',None]),'质量.xlsx',1,'物理')
    assert len(high1['classes'])==2
    walking=parse(quality(),'教学班.xlsx',2,'物理')
    assert [c['class_label'] for c in walking['classes']]==['合成A','合成a']
    admin=parse(quality(2,'administrative',labels=['1','2'],teachers=['合成教师',None]),'质量.xlsx',2,'物理')
    assert len(admin['classes'])==2 and admin['excluded']==[]


def school_detail(rows=None, headers=None, header_row=2):
    return xlsx([['物理学生成绩'], *([[]] * (header_row - 2)),
        headers or ['学号', '学籍', '姓名', '原始', '等级', '年级百分比'],
        *(rows if rows is not None else [['SYN001', 1, '合成甲', 60, 55, '25%']])])


def school_quality(labels=None, teachers=None, avg=50, ratio='83.3%'):
    # 与学校两类原表一致：无“班级类型”列，类别必须由文件名明确。
    w = load_workbook(BytesIO(quality(labels=labels, teachers=teachers, avg=avg)))
    for row in w.active.iter_rows(min_row=3):
        row[2].value = 12  # 应考 12、实考 10，官方比率按应考口径。
        row[18].value = ratio  # C+ 档比率，其他档仍为 0。
    b = BytesIO()
    w.save(b)
    return b.getvalue()


@pytest.mark.parametrize('raw,level,percentile', [
    ('原始', '等级', '年级百分比'),
    ('原始分', '等级分', '年级百分位'),
    ('原始分数', '等级分数', '年级百分位'),
])
def test_school_detail_headers_and_values(raw, level, percentile):
    data = school_detail(headers=['学号', '学籍', '姓名', raw, level, percentile],
        rows=[['00001', 1, '合成甲', 0, 40, '0%'], ['00002', 3, '合成乙', None, None, None]])
    result = parse(data, '明细.xlsx', 2, '物理')
    a, b = result['rows']
    assert a == {'alias': '00001', 'name': '合成甲', 'source_row': 3,
                 'score': 0, 'grade_score': 40, 'grade_percentile': 0}
    assert b['score'] is None and b['grade_score'] is None and b['grade_percentile'] is None
    assert 'xueji_rank' not in a and 'class_kind' not in result
    # 文件名分类仅用于质量表，不误伤明细。
    assert parse(data, '教学班行政班学生.xlsx', 2, '物理')['kind'] == 'student_scores'


@pytest.mark.parametrize('header_row', [2, 10])
def test_school_detail_header_scan_and_original_row_numbers(header_row):
    result = parse(school_detail(header_row=header_row), '明细.xlsx', 2, '物理')
    assert result['rows'][0]['source_row'] == header_row + 1
    with pytest.raises(ValueError, match=f'第 {header_row + 1} 行 原始'):
        parse(school_detail(header_row=header_row, rows=[['SYN001', 1, '合成甲', 'bad', 55, None]]), '明细.xlsx', 2, '物理')


def test_school_detail_scan_limit_duplicate_aliases_and_numeric_id():
    with pytest.raises(ValueError, match='前 10 行'):
        parse(school_detail(header_row=11), '明细.xlsx', 2, '物理')
    for headers in [
        ['学号', '姓名', '原始', '原始分数'],
        ['学号', '姓名', '原始', '等级', '等级分数'],
        ['学号', '姓名', '原始', '年级百分比', '年级百分位'],
    ]:
        with pytest.raises(ValueError, match='重复字段'):
            parse(school_detail(headers=headers), '明细.xlsx', 2, '物理')
    w = load_workbook(BytesIO(school_detail()))
    w.active['A3'] = 1
    w.active['A3'].number_format = '00000'
    b = BytesIO(); w.save(b)
    assert parse(b.getvalue(), '明细.xlsx', 2, '物理')['rows'][0]['alias'] == '00001'


@pytest.mark.parametrize('filename,kind', [
    ('教学班.xlsx', 'teaching'),
    ('2031 教 学 班（副本）.xlsx', 'teaching'),
    ('行政班.xlsx', 'administrative'),
    ('2031 行 政 班导出 (2).xlsx', 'administrative'),
    ('/行政班目录/教学班.xlsx', 'teaching'),
    ('C:\\教学班目录\\行政班.xlsx', 'administrative'),
])
def test_quality_filename_classification_without_class_type(filename, kind):
    r = parse(school_quality(labels=['1', '2'], teachers=['合成教师', None]), filename, 2, '物理')
    assert r['class_kind'] == kind
    assert all(c['class_kind'] == kind for c in r['classes'])
    assert len(r['classes']) == 2
    assert r['warnings'] == []
    assert r['excluded'] == []


def test_quality_filename_priority_fallback_and_ambiguity():
    # 文件名优先，即便存在“班级类型”列也不覆盖明确的教学班文件名。
    admin = quality(kind='administrative', labels=['1', '2'], teachers=['合成教师', None])
    assert len(parse(admin, '教学班.xlsx', 2, '物理')['classes']) == 2
    assert parse(admin, '质量.xlsx', 2, '物理')['class_kind'] == 'administrative'
    for data, name, error in [
        (admin, '教学班行政班.xlsx', '只能包含'),
        (school_quality(), '质量.xlsx', '文件名中添加'),
        (school_quality(), '/教学班目录/质量.xlsx', '文件名中添加'),
    ]:
        with pytest.raises(ValueError, match=error):
            parse(data, name, 2, '物理')
    # 高一仍保留教师空白的真实行政班。
    high1 = parse(quality(1, 'administrative', teachers=[None, None]), '高一质量.xlsx', 1, '物理')
    assert len(high1['classes']) == 2


def test_school_mixed_import_official_values_and_revision(env):
    e = env
    roster = [('SYN001', '合成甲'), ('SYN002', '合成乙')]
    ids = []
    for i, (alias, name) in enumerate(roster):
        tc = TeachingClass(academic_year_id=e.year, subject='物理', label=f'合成教学{i}')
        e.db.add(tc); e.db.commit(); e.tc = tc.id; ids.append(tc.id)
        confirm(e, preview(e, name='合成摸底', files=[('明细.xlsx', detail(rows=[[alias, name, 60, None]]))]))
    before = (e.db.query(TeachingClass).count(), e.db.query(TeachingClassMember).count(), e.db.query(WsStudentIdentity).count())
    rows = [['SYN001', 1, '合成甲', 0, 40, '25%'], ['SYN002', 3, '合成乙', 80, None, None]]
    files = [('明细.xlsx', school_detail(rows)),
        ('教学班.xlsx', school_quality(labels=['合成教学0', '合成教学1'], teachers=['合成教师', None])),
        ('行政班.xlsx', school_quality(labels=['1', '2'], teachers=['合成教师', None]))]
    p = preview(e, files=files, target=False, auto_assign='true')
    assert p['errors'] == [] and all(f['parsed_ok'] and f['warnings'] == [] for f in p['items'])
    assert [f.get('class_kind') for f in p['items']] == [None, 'teaching', 'administrative']
    assert {c['class_kind'] + ':' + c['class_label'] for c in p['roster_changes']} == {
        'teaching:合成教学0', 'teaching:合成教学1', 'administrative:1', 'administrative:2'}
    r = confirm(e, p)
    assert r['imported'] == 6
    assert before == (e.db.query(TeachingClass).count(), e.db.query(TeachingClassMember).count(), e.db.query(WsStudentIdentity).count())
    facts = e.db.query(ScoreFact).filter_by(teaching_exam_id=r['exam_id']).all()
    assert {f.class_ref_id for f in facts} == set(ids)
    assert sorted((f.score, f.grade_score, f.grade_percentile) for f in facts) == [(0, 40, .25), (80, None, None)]
    response = e.client.get(API + f"/exams/{r['exam_id']}/compare")
    assert response.status_code == 200
    assert len(response.json()['classes']) == 4
    # is_managed（2026-10-22 柱状图高亮）：教学班 label 命中或行政班任课教师同名
    managed = {c['class_label']: c['is_managed'] for c in response.json()['classes']}
    assert managed == {'合成教学0': True, '合成教学1': True, '1': True, '2': False}, managed
    for c in response.json()['classes']:
        m = c['metrics']
        assert m['expected_count'] == 12 and m['actual_count'] == 10
        assert c['raw_avg'] == 50 and c['grade_avg'] == 55 and m['stddev'] == 8
        assert next(b for b in m['bands'] if b['label'] == 'C+')['ratio'] == .833
    assert confirm(e, preview(e, files=files, target=False, auto_assign='true'))['skipped'] == 6
    revised = files[:2] + [('行政班.xlsx', school_quality(labels=['1', '2'], teachers=['合成教师', None], ratio='79%'))]
    p = preview(e, files=revised, target=False, auto_assign='true')
    assert p['conflicts']
    confirm(e, p, 409)
    confirm(e, p, revise=True)
    saved = e.db.query(TeachingClassQuality).filter_by(exam_id=r['exam_id'], class_kind='administrative', class_label='1').one()
    assert next(b for b in saved.metrics['bands'] if b['label'] == 'C+')['ratio'] == .79


def test_quality_filename_failure_rolls_back_whole_batch(env):
    e = env
    files = [('明细.xlsx', school_detail()), ('教学班.xlsx', school_quality()), ('质量.xlsx', school_quality())]
    before = (e.db.query(ScoreFact).count(), e.db.query(TeachingClassQuality).count(), e.db.query(WsStudentIdentity).count())
    p = preview(e, files=files)
    assert not p['items'][2]['parsed_ok'] and '文件名中添加' in p['items'][2]['message']
    confirm(e, p, 409)
    assert before == (e.db.query(ScoreFact).count(), e.db.query(TeachingClassQuality).count(), e.db.query(WsStudentIdentity).count())
    assert e.db.query(TeachingExamScope).count() == 0


def test_preview_returns_all_current_name_conflicts_without_writes(env):
    e = env
    confirm(e, preview(e, files=[('明细.xlsx', detail(rows=[
        ['SYN001', '合成甲', 60, None], ['SYN002', '合成乙', 70, None],
    ]))]))
    before = (e.db.query(ScoreFact).count(), e.db.query(ImportBatch).count(), e.db.query(WsStudentIdentity).count())
    r = e.client.post(API + '/imports/preview', data={
        'academic_year_id': str(e.year), 'grade': '2', 'subject': '物理',
        'exam_name': '合成名称核验', 'exam_date': '2030-11-01', 'auto_assign': 'true',
    }, files=[('files', ('明细.xlsx', school_detail(rows=[
        ['SYN001', 1, '合成甲06', 60, 55, None], ['SYN002', 1, '合成乙08', 70, 61, None],
    ]), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))])
    assert r.status_code == 409
    assert r.json()['detail'] == '学号姓名冲突，整批拒绝'
    assert r.json()['conflicts'] == [
        '同学号不同姓名：SYN001（库内「合成甲」≠ 文件「合成甲06」）',
        '同学号不同姓名：SYN002（库内「合成乙」≠ 文件「合成乙08」）',
    ]
    assert before == (e.db.query(ScoreFact).count(), e.db.query(ImportBatch).count(), e.db.query(WsStudentIdentity).count())


def test_initial_supplement_blank_idempotence_and_conflict(env):
    e=env
    p=preview(e);r=confirm(e,p)
    exam=r['exam_id']
    f=e.db.query(ScoreFact).filter_by(teaching_exam_id=exam).one()
    assert f.score==60 and f.grade_score is None
    p=preview(e,files=[('补录.xlsx',detail(rows=[['SYN001','合成甲',60,67]]))],operation='supplement')
    r=confirm(e,p)
    assert r['revised']==1
    f=e.db.query(ScoreFact).filter_by(teaching_exam_id=exam).one()
    assert (f.score,f.grade_score)==(60,67)
    p=preview(e,operation='supplement');r=confirm(e,p)
    assert r['skipped']==1
    assert e.db.query(ScoreFact).filter_by(teaching_exam_id=exam).one().grade_score==67
    p=preview(e,files=[('补录.xlsx',detail(rows=[['SYN001','合成甲',59,64]]))],operation='supplement')
    assert p['errors'][0]['field']=='score'
    confirm(e,p,409,revise=True)
    assert e.db.query(ScoreFact).filter_by(teaching_exam_id=exam).one().score==60
    p=preview(e,files=[('补录.xlsx',detail(rows=[['SYN001','合成甲',60,64]]))],operation='supplement')
    assert p['conflicts'][0]['existing']==67
    confirm(e,p,409)
    confirm(e,p,revise=True)
    rows=e.client.get(API+f'/exams/{exam}/students').json()['students']
    assert len(rows)==1 and rows[0]['grade_label']=='B+'
    assert rows[0]['score']==60
    confirm(e,p,409)


def test_grade1_rank_percentile_patch(env):
    e=env
    exam=confirm(e,preview(e,grade=1))['exam_id']
    p=preview(e,grade=1,operation='supplement',files=[('补录.xlsx',detail(1,[['SYN001','合成甲',None,15,'5.48%']]))])
    confirm(e,p)
    p=preview(e,grade=1,operation='supplement',files=[('空值.xlsx',detail(1,[['SYN001','合成甲',60,None,None]]))])
    confirm(e,p)
    row=e.client.get(API+f'/exams/{exam}/students').json()['students'][0]
    assert row['score']==60 and row['xueji_rank']==15 and row['grade_percentile']==0.0548


def test_quality_only_and_atomic_mixed_population(env):
    e=env
    before=(e.db.query(TeachingClass).count(),e.db.query(TeachingClassMember).count(),e.db.query(WsStudentIdentity).count())
    files=[('教学班.xlsx',quality()),('行政.xlsx',quality(2,'administrative',labels=['1','2'],teachers=['合成教师',None]))]
    p=preview(e,files=files,target=False)
    assert len(p['roster_changes'])==4
    r=confirm(e,p)
    assert before==(e.db.query(TeachingClass).count(),e.db.query(TeachingClassMember).count(),e.db.query(WsStudentIdentity).count())
    cmp=e.client.get(API+f"/exams/{r['exam_id']}/compare").json()
    assert len(cmp['classes'])==4 and cmp['overall']['raw'] is None
    assert all(c['grade_avg']==55 for c in cmp['classes'])
    assert len({c['id'] for c in cmp['classes']})==4
    assert e.client.get(API+f"/exams/{r['exam_id']}/students").json()['students']==[]
    assert e.client.get(API+'/exams',params={'academic_year_id':e.year,'grade':2}).json()['exams'][0]['id']==r['exam_id']
    again=confirm(e,preview(e,files=files,target=False))
    assert again['imported']==0 and again['skipped']==4


def test_grade1_all_classes_and_weighted_official(env):
    e=env
    r=confirm(e,preview(e,grade=1,target=False,files=[('高一质量.xlsx',quality(1,'administrative',labels=['1','2'],teachers=[None,None]))]))
    cmp=e.client.get(API+f"/exams/{r['exam_id']}/compare").json()
    assert len(cmp['classes'])==2
    assert cmp['overall']['raw']['value']==50 and cmp['overall']['grade'] is None
    assert cmp['classes'][0]['raw_diff']==0
    assert cmp['classes'][0]['metrics']['score_sum']==500


def test_scope_and_fact_drift_rolls_back(env):
    e=env
    p=preview(e)
    e.db.get(TeachingClass,e.tc).label='变更后的合成班';e.db.commit()
    confirm(e,p,409)
    assert not e.db.query(TeachingExamScope).filter_by(academic_year_id=e.year).count()
    assert e.db.query(ImportBatch).filter_by(token=p['token']).one().status=='pending'


def test_snapshot_history_quality_revision(env):
    e=env
    files=[('教学班.xlsx',quality(labels=['合成A']))]
    old=confirm(e,preview(e,files=files,target=False))['exam_id']
    p=preview(e,files=[('教学班.xlsx',quality(labels=['合成a'],avg=55))],target=False,name='合成期末')
    confirm(e,p)
    assert len(e.db.get(TeachingExamScope,old).class_snapshot)==1
    assert e.db.query(TeachingComparisonRoster).filter_by(academic_year_id=e.year).count()==2
    p=preview(e,files=[('教学班.xlsx',quality(labels=['合成A'],avg=60))],target=False)
    assert any(c['field']=='raw_avg' for c in p['conflicts'])
    confirm(e,p,409)
    confirm(e,p,revise=True)
    assert e.client.get(API+f'/exams/{old}/compare').json()['classes'][0]['raw_avg']==60


def test_same_name_different_grades_and_subject_guard(env):
    e=env
    a=confirm(e,preview(e))['exam_id']
    tc=TeachingClass(academic_year_id=e.year,subject='物理',label='高一合成班');e.db.add(tc);e.db.commit()
    e.tc=tc.id
    b=confirm(e,preview(e,grade=1,files=[('明细.xlsx',detail(1,[['SYN002','合成乙',70,10,.1]]))]))['exam_id']
    assert a!=b
    assert e.client.get(API+f'/exams/{a}/students').json()['students'][0]['name']=='合成甲'
    assert e.client.get(API+f'/exams/{b}/students').json()['students'][0]['name']=='合成乙'
    r=e.client.get('/api/v1/teaching/analysis/class-compare',params={'exam_name':'合成期中','academic_year_id':e.year})
    assert r.status_code==422
    r=e.client.get(API+'/exams',params={'academic_year_id':e.year,'grade':1,'subject':'化学'})
    assert r.status_code==404


def test_supplement_no_identity_creation_and_partial_failure(env):
    e=env
    before=e.db.query(WsStudentIdentity).count()
    p=preview(e,operation='supplement')
    confirm(e,p,409)
    assert e.db.query(WsStudentIdentity).count()==before
    p=preview(e,files=[('明细.xlsx',detail()),('损坏.xlsx',b'not an xlsx')])
    confirm(e,p,409)
    assert not e.db.query(TeachingExamScope).filter_by(academic_year_id=e.year).count()


def test_transaction_injected_write_failure(env):
    from sqlalchemy import event
    e=env
    p=preview(e,files=[('明细.xlsx',detail()),('教学班.xlsx',quality(labels=['合成A']))])
    def fail(connection,cursor,statement,parameters,context,executemany):
        if statement.startswith('INSERT INTO teaching_class_quality'):
            raise RuntimeError('synthetic transaction failure')
    before=e.db.query(WsStudentIdentity).count()
    event.listen(engine,'before_cursor_execute',fail)
    try:
        with pytest.raises(RuntimeError,match='synthetic transaction failure'):
            e.client.post(API+'/imports/confirm',json={'token':p['token'],'accept_roster_changes':True})
    finally:
        event.remove(engine,'before_cursor_execute',fail)
    e.db.expire_all()
    assert e.db.query(WsStudentIdentity).count()==before
    assert not e.db.query(TeachingExamScope).filter_by(academic_year_id=e.year).count()
    assert e.db.query(ImportBatch).filter_by(token=p['token']).one().status=='pending'
    confirm(e,p)


@pytest.mark.parametrize('value',[True,-1,'NaN','Infinity','wrong',68])
def test_invalid_grade_values(value):
    with pytest.raises(ValueError):
        parse(detail(rows=[['00001','合成甲',0,value]]),'非法.xlsx',2,'物理')


def test_merged_headers_and_summaries():
    # 真正的两行合并表头，空教师不能被向下填充。
    w=Workbook();ws=w.active
    from openpyxl import load_workbook
    ws=load_workbook(BytesIO(quality(1,'administrative',labels=['1','2'],teachers=['合成教师',None]))).active
    ws.merge_cells('D1:E1')
    ws.append(['合成班型小计',None,None,20,20,1000,50,8,90,10,*([0]*8)])
    b=BytesIO();ws.parent.save(b)
    r=parse(b.getvalue(),'合并表头.xlsx',1,'物理')
    assert len(r['classes'])==2 and len(r['aggregates'])==1
    assert r['classes'][1]['teacher_name'] is None


def test_shifted_summary_and_conflicting_totals(env):
    from openpyxl import load_workbook
    w=load_workbook(BytesIO(quality(1,'administrative',labels=['1','2'],teachers=[None,None])))
    ws=w.active
    # 合计行 A:B 合并、C 开始放应考，真实模板的行级位移。
    ws.append(['合成班型合计',None,20,20,1000,50,8,90,10,*([0]*8)])
    ws.merge_cells('A5:B5')
    b=BytesIO();w.save(b)
    r=parse(b.getvalue(),'位移.xlsx',1,'物理')
    a=r['aggregates'][0]
    assert a['metrics']['actual_count']==20 and a['metrics']['raw_avg']==50
    assert a['coverage_verified']
    e=env
    exam=confirm(e,preview(e,grade=1,target=False,files=[('位移.xlsx',b.getvalue())]))['exam_id']
    assert e.client.get(API+f'/exams/{exam}/compare').json()['overall']['raw']['method']=='matching_aggregate'


def test_complementary_files_and_conflicting_duplicate(env):
    e=env
    p=preview(e,grade=1,files=[('分数.xlsx',detail(1,[['00001','合成甲',0,None,None]])),
        ('排名.xlsx',detail(1,[['00001','合成甲',None,20,.2]]))])
    exam=confirm(e,p)['exam_id']
    row=e.client.get(API+f'/exams/{exam}/students').json()['students'][0]
    assert row['score']==0 and row['xueji_rank']==20 and row['grade_percentile']==.2
    fact=e.db.query(ScoreFact).filter_by(teaching_exam_id=exam).one()
    assert fact.field_sources['score']['filename']=='分数.xlsx'
    assert fact.field_sources['xueji_rank']['filename']=='排名.xlsx'
    p=preview(e,grade=1,files=[('甲.xlsx',detail(1,[['00001','合成甲',0,20,.2]])),
        ('乙.xlsx',detail(1,[['00001','合成甲',0,21,.2]]))])
    assert p['errors'][0]['field']=='xueji_rank'
    confirm(e,p,409,revise=True)


def auto(e,rows,name='合成期中',**extra):
    """多班学生合并在一个文件里、按学号自动分班的预览。"""
    return preview(e,files=[('三班合并.xlsx',detail(rows=rows))],target=False,auto_assign='true',name=name,**extra)


def test_auto_assign_splits_multi_class_file(env):
    e=env
    # 三个班各自先单独导入一次摸底，建立学号别名与成员名册
    roster=[('SYN001','合成甲',61),('SYN002','合成乙',62),('SYN003','合成丙',63)]
    tcs=[]
    for i,(sid,sname,score) in enumerate(roster,1):
        tc=TeachingClass(academic_year_id=e.year,subject='物理',label=f'合成{i}班')
        e.db.add(tc);e.db.commit()
        tcs.append(tc.id);e.tc=tc.id
        confirm(e,preview(e,name='合成摸底',files=[('明细.xlsx',detail(rows=[[sid,sname,score,None]]))]))
    rows=[['SYN001','合成甲',71,None],['SYN003','合成丙',83,None],['SYN002','合成乙',72,None]]
    p=auto(e,rows)
    assert p['errors']==[]
    r=confirm(e,p)
    assert r['imported']==3
    alias_ident={a.alias_value:a.identity_id for a in e.db.query(WsStudentAlias).filter_by(data_domain='teaching')}
    facts={f.identity_id:f.class_ref_id for f in e.db.query(ScoreFact).filter_by(exam_name='合成期中')}
    assert [facts[alias_ident[sid]] for sid,_,_ in roster]==tcs
    listed=e.client.get(API+f"/exams/{r['exam_id']}/students").json()['students']
    assert sorted(x['class_label'] for x in listed)==['合成1班','合成2班','合成3班']
    assert confirm(e,auto(e,rows))['skipped']==3
    p=auto(e,[['SYN001','合成甲',71,70],['SYN002','合成乙',72,64]],operation='supplement')
    confirm(e,p)
    grades={f.identity_id:f.grade_score for f in e.db.query(ScoreFact).filter_by(exam_name='合成期中')}
    assert grades[alias_ident['SYN001']]==70 and grades[alias_ident['SYN002']]==64


def test_auto_assign_rejects_unmatched_students(env):
    e=env
    confirm(e,preview(e))  # SYN001 合成甲 → 合成A 班
    ident=e.db.query(WsStudentAlias).filter_by(alias_value='SYN001').one().identity_id
    # 新学号：自动分班不创建学生
    p=auto(e,[['SYN001','合成甲',70,None],['SYN009','合成壬',90,None]])
    errs={x['alias']:x['message'] for x in p['errors'] if 'alias' in x}
    assert '未解析为已有学生' in errs['SYN009']
    confirm(e,p,409)
    # 409 整批零写入：合成期中仍只有预置导入那一条
    assert e.db.query(ScoreFact).filter_by(exam_name='合成期中').count()==1
    assert not e.db.query(WsStudentIdentity).filter_by(display_name='合成壬').count()
    # 有身份但考试当日不在任何班（成员有效期已结束）
    member=e.db.query(TeachingClassMember).filter_by(identity_id=ident).one()
    member.valid_to=date(2030,10,31);e.db.commit()
    errs={x['alias']:x['message'] for x in auto(e,[['SYN001','合成甲',70,None]])['errors']}
    assert '不在你任教的任何教学班名册' in errs['SYN001']
    # 同一学号同时是两个班的成员
    member.valid_to=None;e.db.commit()
    other=TeachingClass(academic_year_id=e.year,subject='物理',label='合成B班')
    e.db.add(other);e.db.flush()
    e.db.add(TeachingClassMember(teaching_class_id=other.id,identity_id=ident,valid_from=date(2030,9,1)));e.db.commit()
    errs={x['alias']:x['message'] for x in auto(e,[['SYN001','合成甲',70,None]])['errors']}
    assert '同时属于' in errs['SYN001']
    # 指定班级与自动分班互斥
    form={'academic_year_id':str(e.year),'grade':'2','subject':'物理','exam_name':'合成期中',
        'exam_date':'2030-11-01','operation':'initial','teaching_class_id':str(e.tc),'auto_assign':'true'}
    r=e.client.post(API+'/imports/preview',data=form,files=[('files',('明细.xlsx',detail(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))])
    assert r.status_code==422


def test_history_candidates_and_own_domain_only(env):
    e=env
    first=confirm(e,preview(e))['exam_id']
    old_person=e.db.query(ScoreFact).filter_by(teaching_exam_id=first).one().identity_id
    year=AcademicYear(name='后续合成学年',start_date=date(2030,9,1),end_date=date(2031,8,31))
    e.db.add(year);e.db.flush()
    tc=TeachingClass(academic_year_id=year.id,subject='物理',label='合成新班');e.db.add(tc);e.db.commit()
    e.year,e.tc=year.id,tc.id
    p=preview(e)
    assert p['identity_candidates'][0]['person_id']==old_person
    confirm(e,p,409)
    confirm(e,p,identity_confirmations={'SYN001':old_person})
    assert e.db.query(ScoreFact).filter_by(academic_year_id=year.id).one().identity_id==old_person


def test_new_division_replaces_only_future_snapshot(env):
    e=env
    old=confirm(e,preview(e,target=False,files=[('旧教学班.xlsx',quality())]))['exam_id']
    p=preview(e,target=False,name='合成期末',exam_date='2031-01-10',replace_roster='true',
        files=[('新教学班.xlsx',quality(labels=['合成b']))])
    assert [c['action'] for c in p['roster_changes']].count('remove')==2
    current=confirm(e,p)['exam_id']
    assert len(e.db.get(TeachingExamScope,old).class_snapshot)==2
    assert [c['class_label'] for c in e.db.get(TeachingExamScope,current).class_snapshot]==['合成b']
    # 修正旧考试均分不能把最新分班参照倒退回旧班名。
    confirm(e,preview(e,target=False,files=[('修正教学班.xlsx',quality(avg=60))]),revise=True)
    latest=e.db.query(TeachingComparisonRoster).filter_by(academic_year_id=e.year).order_by(TeachingComparisonRoster.version.desc()).first()
    assert [c['class_label'] for c in latest.classes]==['合成b']


def test_numeric_id_display_zeroes_preserved():
    from openpyxl import load_workbook
    w=load_workbook(BytesIO(detail(rows=[[1,'合成甲',0,None]])))
    w.active['A2'].number_format='00000'
    b=BytesIO();w.save(b)
    assert parse(b.getvalue(),'前导零.xlsx',2,'物理')['rows'][0]['alias']=='00001'


def test_legacy_score_date_conflict_does_not_attach_wrong_exam(env):
    e=env
    first=confirm(e,preview(e))['exam_id']
    fact=e.db.query(ScoreFact).filter_by(teaching_exam_id=first).one()
    fact.teaching_exam_id=None
    e.db.commit()
    e.db.delete(e.db.get(TeachingExamScope,first));e.db.commit()
    p=preview(e,operation='supplement',exam_date='2030-11-02',files=[('补录.xlsx',detail(rows=[['SYN001','合成甲',60,67]]))])
    assert p['exam_resolution']['requires_choice']
    assert any('补录须先存在' in err['message'] for err in p['errors'])
    confirm(e,p,409,revise=True)
    assert not e.db.query(TeachingExamScope).count()
    assert e.db.query(ScoreFact).one().grade_score is None
    r=confirm(e,preview(e,operation='supplement',files=[('补录.xlsx',detail(rows=[['SYN001','合成甲',60,67]]))]))
    assert e.db.query(ScoreFact).filter_by(teaching_exam_id=r['exam_id']).one().grade_score==67


def test_ai_uses_same_official_service(env):
    from app.api.chat_tools import resolve_scope_snapshot,execute_session_tool
    e=env
    exam=confirm(e,preview(e,target=False,files=[('教学班.xlsx',quality())]))['exam_id']
    snap=resolve_scope_snapshot(e.db,1,'teaching',academic_year_id=e.year,teaching_class_id=e.tc,subject='物理')
    result=execute_session_tool(e.db,snap,'get_class_comparison',{'exam_id':exam})
    http=e.client.get(API+f'/exams/{exam}/compare').json()
    assert result==http
    catalog=execute_session_tool(e.db,snap,'get_exam_list',{})
    assert catalog['comparison_exams'][0]['exam_id']==exam


def test_migration_roundtrip_preserves_scores(tmp_path):
    import os,subprocess,sys
    from pathlib import Path
    root=Path(__file__).resolve().parents[2]
    data=tmp_path/'data';data.mkdir()
    env={**os.environ,'EXAM_TRACKER_DIR':str(data),'EXAM_TRACKER_BACKUP_DIR':str(tmp_path/'backups'),'PYTHONPATH':str(root)}
    source='''
from app.db.schema import ensure_app_schema,_alembic_config
from alembic import command
from app.db.models import engine
from sqlalchemy import text
ensure_app_schema()
cfg=_alembic_config()
command.downgrade(cfg,'0018')
with engine.begin() as c:
 c.execute(text("INSERT INTO academic_year(id,name,start_date,end_date) VALUES(1,'合成迁移学年','2030-09-01','2031-08-31')"))
 c.execute(text("INSERT INTO ws_student_identity(id,data_domain,display_name) VALUES(1,'teaching','合成甲')"))
 c.execute(text("INSERT INTO score_fact(data_domain,academic_year_id,exam_name,identity_id,subject,score,subject_key,total_key,source,data_revision,exam_date_precision) VALUES('teaching',1,'合成考试',1,'物理',0,'物理','','synthetic',1,'day')"))
command.upgrade(cfg,'head')
with engine.connect() as c:
 assert c.execute(text('SELECT score,teaching_exam_id FROM score_fact')).one()==(0,None)
 assert c.execute(text('PRAGMA foreign_key_check')).fetchall()==[]
 # 2026-10 迁移链头已推进到 0021（rank band 导入与 xueji 列），升级终点随链更新。
 assert c.execute(text('SELECT version_num FROM alembic_version')).scalar()=='0021'
'''
    r=subprocess.run([sys.executable,'-c',source],cwd=root,env=env,capture_output=True,text=True)
    assert r.returncode==0,r.stdout+r.stderr


def test_canonical_event_reuse_and_same_name_different_date_choice(env):
    from app.db.workspace_models import WorkspaceExamEvent
    e=env
    p=preview(e)
    assert p['exam_resolution']['selected'] is None and not p['exam_resolution']['requires_choice']
    first=confirm(e,p)['exam_id']
    event=e.db.get(TeachingExamScope,first).exam_event_id
    p=preview(e)
    assert p['exam_resolution']['selected']==event and not p['exam_resolution']['requires_choice']
    assert confirm(e,p)['skipped']==1
    p=preview(e,exam_date='2030-12-01')
    assert p['exam_resolution']['requires_choice'] and p['exam_resolution']['selected'] is None
    confirm(e,p,409)
    assert e.db.query(WorkspaceExamEvent).count()==1
    second=confirm(e,p,exam_choice='new')['exam_id']
    assert second!=first and e.db.query(TeachingExamScope).count()==2
    for eid in (first,second):
        assert e.client.get(API+f'/exams/{eid}/students').status_code==200
        assert e.client.get(API+f'/exams/{eid}/compare').status_code==200
    assert e.db.query(ScoreFact).count()==2
    for member in e.db.query(TeachingClassMember):
        member.valid_from=date(2025,9,1)
    e.db.commit()
    keys=e.client.get('/api/v1/shared/exams',params={'mode':'teaching','academic_year_id':e.year,'teaching_class_id':e.tc}).json()['exams']
    assert len({r['exam_key'] for r in keys})==2
    for r in keys:
        response=e.client.get('/api/v1/teaching/analysis/exams/'+r['exam_key']+'/students',params={'academic_year_id':e.year,'teaching_class_id':e.tc})
        assert response.status_code==200 and len(response.json()['students'])==1,response.text
    assert e.client.get('/api/v1/teaching/analysis/exams/合成期中/students',params={'academic_year_id':e.year,'teaching_class_id':e.tc}).status_code==422


def test_event_candidate_drift_rejects_without_scores(env):
    from app.db.workspace_models import WorkspaceExamEvent
    e=env
    stale=preview(e)
    first=confirm(e,preview(e))
    original=e.db.query(ScoreFact).count()
    confirm(e,stale,409)
    assert e.db.query(ScoreFact).count()==original and e.db.query(WorkspaceExamEvent).count()==1


def full_school_quality(zero=False, chinese=91.2):
    """学校同构 33 列，新增四项均分。仅合成值，无真实学生/教师。"""
    from openpyxl.utils import get_column_letter
    bands=list(GRADE_LABELS.values())
    top=['班级','任课教师','人数',None,'平均分',None,'标准差','语文','数学','英语','三大门','等级人数及比率']+[None]*21
    sub=[None,None,'应考','实考','原始','等级',None,None,None,None,None]
    for band in bands:sub += [band,'比率']
    row=['02',None,0 if zero else 12,0 if zero else 10,0 if zero else 65.5,
         0 if zero else 55,0 if zero else 8,chinese,0,None,280.7]
    for band in bands:row += [0 if zero else 10 if band=='C+' else 0,'0%' if zero or band!='C+' else '83.3%']
    return xlsx([top,sub,row], ['A1:A2','B1:B2','C1:D1','E1:F1','L1:AG1']+
        [f'{get_column_letter(i)}1:{get_column_letter(i)}2' for i in range(7,12)])


@pytest.mark.parametrize('kind',['行政班','教学班'])
@pytest.mark.parametrize('zero',[False,True])
def test_complete_quality_metrics_and_empty_teacher_persist(env,kind,zero):
    e=env
    before=(e.db.query(TeachingClass).count(),e.db.query(TeachingClassMember).count(),e.db.query(WsStudentIdentity).count())
    files=[(kind+'.xlsx',full_school_quality(zero))]
    p=preview(e,files=files,target=False)
    assert len(p['items'][0]['classes'])==1 and not p['items'][0]['excluded']
    r=confirm(e,p)
    c=e.client.get(API+f"/exams/{r['exam_id']}/compare").json()['classes'][0]
    m=c['metrics']
    assert c['teacher_name'] is None
    assert (m['chinese_avg'],m['math_avg'],m['english_avg'],m['main_three_avg'])==(91.2,0,None,280.7)
    assert m['actual_count']==(0 if zero else 10)
    assert next(b for b in m['bands'] if b['label']=='C+')['ratio']==(0 if zero else .833)
    assert confirm(e,preview(e,files=files,target=False))['skipped']==1
    p=preview(e,files=[(kind+'.xlsx',full_school_quality(zero,chinese=92))],target=False)
    confirm(e,p,status=409)
    assert e.db.query(TeachingClassQuality).one().metrics['chinese_avg']==91.2
    confirm(e,p,revise=True)
    assert e.db.query(TeachingClassQuality).one().metrics['chinese_avg']==92
    assert before==(e.db.query(TeachingClass).count(),e.db.query(TeachingClassMember).count(),e.db.query(WsStudentIdentity).count())
    assert e.db.query(ScoreFact).count()==0


def test_analysis_details_and_profile_keep_source_percentile_and_scope_rank(env, monkeypatch):
    class FrozenDate(date):
        @classmethod
        def today(cls): return cls(2030, 11, 1)
    monkeypatch.setattr("app.core.context.date", FrozenDate)
    e=env
    # 两班同场，原始分并列/0/缺考；百分比独立于等级分和范围名次。
    first=confirm(e,preview(e,files=[('明细.xlsx',school_detail(rows=[
        ['SYN001',1,'合成甲',90,67,.125],['SYN002',1,'合成乙',None,55,.5],
        ['SYN003',1,'合成丙',0,40,1]]))]))
    tc1=e.tc
    tc2=TeachingClass(academic_year_id=e.year,grade=2,subject='物理',label='合成B')
    e.db.add(tc2);e.db.commit();e.tc=tc2.id
    confirm(e,preview(e,files=[('明细.xlsx',school_detail(rows=[['SYN004',1,'合成丁',90,70,0]]))]))
    from urllib.parse import quote
    exam=e.db.get(TeachingExamScope,first['exam_id'])
    url='/api/v1/teaching/analysis/exams/'+quote('event:'+str(exam.exam_event_id))+'/students'
    params={'academic_year_id':e.year}
    response=e.client.get(url,params=params)
    assert response.status_code==200,response.text
    data=response.json();rows=data['students']
    assert data['grade']==2 and len(rows)==4
    assert [r['rank'] for r in rows]==[1,1,3,None]
    assert {r['name']:r['grade_percentile'] for r in rows}=={'合成甲':.125,'合成乙':.5,'合成丙':1,'合成丁':0}
    assert {r['name']:r['grade_label'] for r in rows}['合成甲']=='A'
    assert all(r['source_domain']=='teaching' and 'totals' not in r for r in rows)
    one=e.client.get(url,params={**params,'teaching_class_id':tc1}).json()['students']
    assert [r['rank'] for r in one]==[1,2,None] and len(one)==3
    person=next(r['person_id'] for r in rows if r['name']=='合成甲')
    prof=e.client.get(f'/api/v1/teaching/students/{person}',params={**params,'teaching_class_id':tc1})
    assert prof.status_code==200,prof.text
    assert prof.json()['subjects'][0]['exams'][0]['grade_percentile']==.125
    # 班主任侧不能读取教学身份，参照不授权学生。
    denied=e.client.get(f'/api/v1/homeroom/students/{person}',params=params)
    assert denied.status_code in (404,409)


def test_analysis_grade1_and_unknown_grade_fields(env):
    e=env
    r=confirm(e,preview(e,grade=1,files=[('明细.xlsx',detail(1,rows=[['SYN001','合成甲',0,10,.2]]))]))
    exam=e.db.get(TeachingExamScope,r['exam_id'])
    url=f'/api/v1/teaching/analysis/exams/event:{exam.exam_event_id}/students'
    data=e.client.get(url,params={'academic_year_id':e.year,'teaching_class_id':e.tc}).json()
    assert data['grade']==1
    assert data['students'][0]['xueji_rank']==10 and data['students'][0]['grade_percentile']==.2
    from app.db.workspace_models import WorkspaceExamEvent
    exam.grade=2  # 显式冲突时未知，不从分数/班名猜。
    e.db.get(WorkspaceExamEvent,exam.exam_event_id).grade=1
    e.db.commit()
    data=e.client.get(url,params={'academic_year_id':e.year,'teaching_class_id':e.tc}).json()
    assert data['grade'] is None
