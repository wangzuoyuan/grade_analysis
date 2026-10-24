"""名次按学籍代码选列 + 全员参与名次统计（2026-10-24 用户裁决，替代
2026-10-20 的「仅统计学籍=1 人口过滤」口径）。

有效名次（definitions.resolve_year_rank 第三参 xueji）：
- 学籍=1 → 学籍排名（xueji_rank），缺时用年级排名（grade_rank）补；
- 学籍为已知非 1 整数（2/3/4 借读类）→ 只认年级排名，绝不回退学籍排名
  （学校旧表给这类学生填的是小圈子名次，非全年级口径）；
- 学籍未采集（NULL）→ 维持历史行为：学籍排名优先、缺则年级排名；
- 两列皆缺 → None（不落段、不进名次区间、频次不计数）。

排名分段/频次/区间/段位分布不再按学籍过滤人口——由有效名次本身决定
谁进统计。种子：甲=1、乙=3、丙=NULL。
"""

from datetime import date
from pathlib import Path
from urllib.parse import quote

import pytest

API = "/api/v1"
EXAM = "P3XJ学籍选列考"
EXAM_FALLBACK = "P3XJ学籍回退考"
EXAM_NO_RANK = "P3XJ无名次考"

_HEADERS = {
    1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
    5: "语文", 6: "数学", 7: "英语",
    20: "+3总分", 21: "主三门总分", 23: "主三门学籍排名",
}


def _write_detail(path: Path, rows) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col, text in _HEADERS.items():
        ws.cell(row=2, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def _upload(client, path: Path, extra: dict):
    return client.post(
        f"{API}/imports/preview",
        data={"mode": "homeroom", **{k: str(v) for k, v in extra.items()}},
        files=[("files", (path.name, path.read_bytes(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
    )


@pytest.fixture(scope="module")
def rank_seed(v1_seed):
    """甲=1（学籍排名 50）、乙=3（假学籍排名 300 / 年级排名 310）、
    丙=NULL（学籍排名 100，未采集维持学籍排名优先）。"""
    from app.db import workspace_models as wm
    from app.db.models import SessionLocal

    s = v1_seed
    db = SessionLocal()

    def set_xueji(ident_id, value):
        ident = db.get(wm.WsStudentIdentity, ident_id)
        ident.xueji = value

    set_xueji(s.jia_h_id, 1)
    set_xueji(s.yi_h_id, 3)
    set_xueji(s.bing_h_id, None)  # 未采集

    def fact(ident_id, exam, score, xueji_rank=None, grade_rank=None,
             subject=None, percentile=None, total_type="主三门"):
        f = wm.ScoreFact()
        f.data_domain = "homeroom"
        f.academic_year_id = s.ay_id
        f.exam_name = exam
        f.exam_date = date(2025, 12, 10)
        f.class_ref_id = s.h6_id
        f.identity_id = ident_id
        f.subject = subject
        f.total_type = total_type
        f.score = score
        f.xueji_rank = xueji_rank
        f.grade_rank = grade_rank
        f.grade_percentile = percentile
        f.source = "p3-xueji-test"
        db.add(f)

    # 主场：甲 50（学籍排名）、乙 300/310（学籍=3 → 有效名次 310）、丙 100（未采集 → 100）
    fact(s.jia_h_id, EXAM, 275.0, xueji_rank=50)
    fact(s.yi_h_id, EXAM, 236.0, xueji_rank=300, grade_rank=310)
    fact(s.bing_h_id, EXAM, 207.0, xueji_rank=100)
    # 回退场：甲=1 但学籍排名缺失、年级排名 55 → 55；乙无年级排名 → 无有效名次
    fact(s.jia_h_id, EXAM_FALLBACK, 275.0, grade_rank=55)
    fact(s.yi_h_id, EXAM_FALLBACK, 236.0, xueji_rank=300)
    # 语文单科行（频次百分位口径 + 区间前%口径用）
    fact(s.jia_h_id, EXAM, 88.0, xueji_rank=10, subject="语文", percentile=0.9, total_type=None)
    fact(s.yi_h_id, EXAM, 76.0, xueji_rank=20, subject="语文", percentile=0.5, total_type=None)
    fact(s.bing_h_id, EXAM, 65.0, xueji_rank=30, subject="语文", percentile=0.3, total_type=None)
    # 无名次考试（三人均无任何名次 → bands 409）
    fact(s.jia_h_id, EXAM_NO_RANK, 275.0)

    db.commit()
    yield
    db.close()


def _q(client, path, **params):
    return client.get(path, params=params)


def test_bands_and_distribution_count_all_members(client, rank_seed, v1_seed):
    s = v1_seed
    # 段位分布（阈值 高分1-80/临界400-500/薄弱501+）：甲 50 高分段；
    # 乙 310、丙 100 均不落段——全员参与统计、各取有效名次
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name=EXAM, subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    body = r.json()
    labels = [b["label"].split("（")[0] for b in body["bands"]]
    assert labels == ["高分段", "临界段", "薄弱段"]
    counts = {b["label"].split("（")[0]: b["count"] for b in body["bands"]}
    assert counts["高分段"] == 1  # 甲（学籍=1，学籍排名 50）
    top = body["bands"][0]
    assert top["students"] == [s.jia_h_id]
    assert top["names"] == ["秦甲"]
    assert all(b["students"] == [] for b in body["bands"][1:])

    # 排名分段：40 名一档，三人各取有效名次落档
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-distribution",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    main3 = next(se for se in r.json()["series"] if se["total_type"] == "主三门")
    assert main3["counts"].get("r41_80") == 1    # 甲 50
    assert main3["counts"].get("r81_120") == 1   # 丙 100（未采集 → 学籍排名）
    assert main3["counts"].get("r281_320") == 1  # 乙 310（学籍=3 → 年级排名，非 300）
    assert sum(main3["counts"].values()) == 3


def test_frequency_and_range_use_effective_rank(client, rank_seed, v1_seed):
    s = v1_seed
    # 频次（总分名次口径）：三人各落一档
    r = _q(client, f"{API}/homeroom/analysis/rank-frequency",
           metric="total:主三门", exam_names=EXAM,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    names = sorted(st["name"] for st in r.json()["students"])
    assert names == ["秦丙", "秦乙", "秦甲"]

    # 区间（总分 1-500）：乙的 year_rank/metric_value=310（年级排名），不是 300
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-range",
           metric="total:主三门", rank_min=1, rank_max=500,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    by_name = {st["name"]: st for st in r.json()["students"]}
    assert set(by_name) == {"秦甲", "秦乙", "秦丙"}
    assert by_name["秦乙"]["year_rank"] == 310
    assert by_name["秦乙"]["metric_value"] == 310
    assert by_name["秦甲"]["year_rank"] == 50
    assert by_name["秦丙"]["year_rank"] == 100

    # 单科口径照旧全部学生（百分位）
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-range",
           metric="subject:语文", rank_min=0, rank_max=100,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    names = sorted(st["name"] for st in r.json()["students"])
    assert names == ["秦丙", "秦乙", "秦甲"]


def test_xueji1_falls_back_to_grade_rank(client, rank_seed, v1_seed):
    """学籍=1 且学籍排名缺失 → 用年级排名补；学籍=3 无年级排名 → 不进统计。"""
    s = v1_seed
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name=EXAM_FALLBACK, subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    counts = {b["label"].split("（")[0]: b["count"] for b in r.json()["bands"]}
    assert counts["高分段"] == 1  # 甲 55（年级排名回退）
    assert sum(counts.values()) == 1  # 乙学籍排名 300 不被取用

    r = _q(client, f"{API}/homeroom/analysis/rank-frequency",
           metric="total:主三门", exam_names=EXAM_FALLBACK,
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    assert [st["name"] for st in r.json()["students"]] == ["秦甲"]


def test_bands_409_when_no_effective_rank(client, rank_seed, v1_seed):
    """全无名次 → 409（F10 语义：有效名次人口为空）。"""
    s = v1_seed
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name=EXAM_NO_RANK, subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 409

    # 学籍更新即时联动：甲改为学籍 4 且只有学籍排名（无年级排名）→ 无有效名次
    from app.db import workspace_models as wm
    from app.db.models import SessionLocal

    db = SessionLocal()
    db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 4
    db.commit()
    db.close()
    try:
        r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/rank-distribution",
               academic_year_id=s.ay_id, class_id=s.h6_id)
        assert r.status_code == 200
        main3 = next(se for se in r.json()["series"] if se["total_type"] == "主三门")
        # 甲退出（学籍=4 只认年级排名，而其只有学籍排名）；乙丙仍在
        assert sum(main3["counts"].values()) == 2
    finally:
        db = SessionLocal()
        db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 1
        db.commit()
        db.close()


def test_students_rows_carry_effective_rank(client, rank_seed, v1_seed):
    """明细表学生行 ranks 下发有效名次（乙=310 非 300）。"""
    s = v1_seed
    r = _q(client, f"{API}/homeroom/analysis/exams/{quote(EXAM)}/students",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 200, r.text
    ranks = {st["name"]: st["ranks"].get("主三门") for st in r.json()["students"]}
    assert ranks.get("秦甲") == 50
    assert ranks.get("秦乙") == 310
    assert ranks.get("秦丙") == 100


def test_import_updates_xueji_from_file(client, v1_seed, tmp_path):
    """明细表导入：文件学籍列有值即覆盖身份 xueji（既有值也会被刷新）。"""
    from app.db import workspace_models as wm
    from app.db.models import SessionLocal
    from tests.v1.conftest import ALIAS_JIA_H

    s = v1_seed
    db = SessionLocal()
    db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 1
    db.commit()
    db.close()

    detail = _write_detail(tmp_path / "学籍更新明细.xlsx", [
        {1: ALIAS_JIA_H, 2: 6, 3: 3, 4: "秦甲", 5: 90, 20: 176, 21: 276, 23: 40},  # 文件学籍=3
    ])
    preview = _upload(client, detail, {
        "class_id": s.h6_id, "exam_name": "P3XJ学籍更新考", "exam_date": "2025-12-11",
    })
    assert preview.status_code == 200, preview.text
    confirmed = client.post(f"{API}/imports/confirm", json={"token": preview.json()["token"]})
    assert confirmed.status_code == 200, confirmed.text

    db = SessionLocal()
    assert db.get(wm.WsStudentIdentity, s.jia_h_id).xueji == 3  # 覆盖为文件值
    db.close()
    # 学籍=3 只认年级排名：新考试文件只有学籍排名 → bands 409
    r = _q(client, f"{API}/homeroom/analysis/bands",
           exam_name="P3XJ学籍更新考", subject="主三门", metric="total",
           academic_year_id=s.ay_id, class_id=s.h6_id)
    assert r.status_code == 409

    db = SessionLocal()
    db.get(wm.WsStudentIdentity, s.jia_h_id).xueji = 1
    db.commit()
    db.close()
