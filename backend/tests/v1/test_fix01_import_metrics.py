"""审查问题 1 回归：旧模板导入完整保留年级百分位/学籍名次/年级名次。

背景（独立代码审查报告 §1 / 核实说明 §1 / probe.py 1_import_metrics）：
H/T 解析器与 imports.py 写入端（inserts/revises）均已支持三指标，但
ws_parse.py 转换层构造 rows 时只拷贝六键，grade_percentile/xueji_rank/
grade_rank 被静默丢弃 → 经 /api/v1/imports 导入的 ScoreFact 三指标恒
NULL。修复后转换层三键透传（值可为 None，None 保持 None 不转 0——
原始缺考/未导入显示「—」是项目铁律）。

指标挂载位置（excel_parser / teaching_excel 输出结构，先读确认）：
- 总分行（仅 homeroom）：三键齐（高二固定列布局无年级排名列，
  H 解析器 grade23 的 grade_rank 恒 None 属既有语义；高一表才有，
  见 tests/test_excel_parser.py::test_parse_grade1_percentiles）；
- 学科行（两域同构）：只有 grade_percentile，xueji_rank/grade_rank
  恒 None——单科本就没有学籍/年级名次列。

验收（docs/planning/2026-10-18-review-remediation-plan.md 第 1 行）：
合成 Excel 指标与入库值一致；空值/缺考/零分保持原义；重复导入同值
幂等 skipped；revise 修订路径保留并更新指标；teaching 模式（高一表
形态，物理带年级百分位）透传不回归。教学平台导入链路
（teaching_platform.py，不经 ws_parse）另由 test_teaching_platform.py
全量回归佐证。
"""

from pathlib import Path

from tests.v1.conftest import ALIAS_BING_H, ALIAS_JIA_H, ALIAS_WU_T, ALIAS_YI_H
from tests.v1.test_f05_f06_f07 import _student_row, _upload, _write_grade23_xlsx

API = "/api/v1"

# 列号含义（_student_row / grade23 固定列布局）：
# 5-7 语数英原始分 | 8 物理 9 物理等级分 | 20 +3总分
# 21 主三门总分 22 主三门百分位 23 主三门学籍排名
# 24 3+3总分 25 3+3百分位 26 3+3学籍排名 | 27-29 语数英百分位
EXAM_ROUNDTRIP = "FIX01指标回填考"
EXAM_IDEMPOTENT = "FIX01幂等考"
EXAM_REVISE = "FIX01修订考"
EXAM_TEACHING = "FIX01教学考"
EXAM_DATE = "2025-11-10"


def _db():
    from app.db.models import SessionLocal

    return SessionLocal()


def _upload_homeroom(client, path, exam_name):
    prev = _upload(
        client,
        path,
        "homeroom",
        {"exam_name": exam_name, "exam_date": EXAM_DATE},
    )
    assert prev.status_code == 200, prev.text
    return prev


def _confirm(client, prev, **extra):
    ok = client.post(
        f"{API}/imports/confirm", json={"token": prev.json()["token"], **extra}
    )
    assert ok.status_code == 200, ok.text
    return ok.json()


def _homeroom_facts(exam_name):
    from app.db import workspace_models as wm

    db = _db()
    try:
        return (
            db.query(wm.ScoreFact)
            .filter_by(exam_name=exam_name, data_domain="homeroom")
            .all()
        )
    finally:
        db.close()


def _metric_rows():
    """三名已知学生（种子 alias）：甲全指标；乙零分/无指标；丙缺考带指标。"""
    jia = _student_row(
        ALIAS_JIA_H, "秦甲", chinese=88, math=92, english=95, physics=90
    )
    jia.update(
        {21: 275, 22: 0.2, 23: 41, 24: 436, 25: 0.3, 26: 15,
         27: 0.98, 28: 0.9, 29: 0.85}
    )
    # 乙：英语零分（0 是有效分不转 None/NULL），全行无任何指标
    yi = _student_row(ALIAS_YI_H, "秦乙", chinese=76, math=81, english=0)
    # 丙：语文缺考（无分）但百分位有值 → 行保留 score=None；主三门仅
    # 学籍排名有值 → score/pct 均为 None 保真
    bing = _student_row(ALIAS_BING_H, "秦丙")
    bing.update({27: 50, 23: 99})
    return [jia, yi, bing]


def test_homeroom_metrics_roundtrip_and_null_semantics(client, v1_seed, tmp_path):
    """总分行三指标 + 学科行百分位全链路透传；None/零分/缺考保真。"""
    path = _write_grade23_xlsx(tmp_path / "FIX01高二明细.xlsx", _metric_rows())
    prev = _upload_homeroom(client, path, EXAM_ROUNDTRIP)
    result = _confirm(client, prev)
    # 甲 4 科 + 2 总分；乙 3 科；丙 语文 + 主三门 = 11 行
    assert result["imported"] == 11, result

    facts = _homeroom_facts(EXAM_ROUNDTRIP)
    assert len(facts) == 11
    by_key = {(f.identity_id, f.subject, f.total_type): f for f in facts}

    # 总分行：score + 百分位 + 学籍名次透传；grade23 无年级排名列 → None
    jia_m3 = by_key[(v1_seed.jia_h_id, None, "主三门")]
    assert (jia_m3.score, jia_m3.grade_percentile, jia_m3.xueji_rank,
            jia_m3.grade_rank) == (275.0, 0.2, 41, None)
    jia_33 = by_key[(v1_seed.jia_h_id, None, "3+3")]
    assert (jia_33.score, jia_33.grade_percentile, jia_33.xueji_rank) == (
        436.0, 0.3, 15
    )
    # 学科行：只有年级百分位（解析器语义），名次键恒 None 不推导
    jia_yw = by_key[(v1_seed.jia_h_id, "语文", None)]
    assert (jia_yw.score, jia_yw.grade_percentile, jia_yw.xueji_rank,
            jia_yw.grade_rank) == (88.0, 0.98, None, None)
    # 乙：零分保持 0.0；无指标学科行三指标 None
    yi_en = by_key[(v1_seed.yi_h_id, "英语", None)]
    assert (yi_en.score, yi_en.grade_percentile, yi_en.xueji_rank) == (0.0, None, None)
    # 丙：缺考行 score NULL + 百分位仍入库；总分缺考但名次保真
    bing_yw = by_key[(v1_seed.bing_h_id, "语文", None)]
    assert (bing_yw.score, bing_yw.grade_percentile) == (None, 0.5)
    bing_m3 = by_key[(v1_seed.bing_h_id, None, "主三门")]
    assert (bing_m3.score, bing_m3.grade_percentile, bing_m3.xueji_rank) == (
        None, None, 99
    )


def test_reimport_same_values_idempotent(client, v1_seed, tmp_path):
    """同值重导：全部 skipped 幂等，不重复写入、指标值不变。"""
    path = _write_grade23_xlsx(tmp_path / "FIX01幂等明细.xlsx", _metric_rows())
    prev = _upload_homeroom(client, path, EXAM_IDEMPOTENT)
    first = _confirm(client, prev)
    assert first["imported"] == 11, first

    again = _upload_homeroom(client, path, EXAM_IDEMPOTENT)
    second = _confirm(client, again)
    assert second["imported"] == 0, second
    assert second["skipped"] == 11, second
    assert second["revised"] == 0, second

    facts = _homeroom_facts(EXAM_IDEMPOTENT)
    assert len(facts) == 11  # 无重复行
    jia_m3 = next(f for f in facts if f.total_type == "主三门")
    assert (jia_m3.score, jia_m3.grade_percentile, jia_m3.xueji_rank) == (
        275.0, 0.2, 41
    )


def test_revise_updates_and_keeps_metrics(client, v1_seed, tmp_path):
    """revise 修订：异值行走修订更新三指标（不清空），未变行 skipped。"""
    path = _write_grade23_xlsx(tmp_path / "FIX01修订初版.xlsx", _metric_rows())
    prev = _upload_homeroom(client, path, EXAM_REVISE)
    assert _confirm(client, prev)["imported"] == 11

    # 甲：五处指标/分数变化（语数英百分位 + 主三门 + 3+3），物理不变
    rows = _metric_rows()
    rows[0].update(
        {21: 280, 22: 0.15, 23: 35, 24: 440, 25: 0.25, 26: 12,
         27: 0.99, 28: 0.91, 29: 0.86}
    )
    revised_path = _write_grade23_xlsx(
        tmp_path / "FIX01修订新版.xlsx", rows
    )
    prev2 = _upload_homeroom(client, revised_path, EXAM_REVISE)
    # 不带 revise：异值整批 409 拒绝、零写入
    denied = client.post(
        f"{API}/imports/confirm", json={"token": prev2.json()["token"]}
    )
    assert denied.status_code == 409, denied.text

    prev3 = _upload_homeroom(client, revised_path, EXAM_REVISE)
    revised = _confirm(client, prev3, revise=True)
    assert revised["revised"] == 5, revised
    assert revised["skipped"] == 6, revised
    assert revised["imported"] == 0, revised

    facts = _homeroom_facts(EXAM_REVISE)
    assert len(facts) == 11  # 修订原地更新，行数不变
    by_key = {(f.identity_id, f.subject, f.total_type): f for f in facts}
    jia_m3 = by_key[(v1_seed.jia_h_id, None, "主三门")]
    assert (jia_m3.score, jia_m3.grade_percentile, jia_m3.xueji_rank) == (
        280.0, 0.15, 35
    )
    jia_yw = by_key[(v1_seed.jia_h_id, "语文", None)]
    assert (jia_yw.score, jia_yw.grade_percentile) == (88.0, 0.99)
    # 未变行：物理仍无指标且分数不变
    jia_wl = by_key[(v1_seed.jia_h_id, "物理", None)]
    assert (jia_wl.score, jia_wl.grade_percentile) == (90.0, None)


def _write_grade1_teaching_xlsx(path: Path, rows) -> Path:
    """高一两行表头教学明细：语/物各带 分数+年级百分位，含教学班标签列。

    T 解析器按表头形态识别高一表（第 2 行 1-4 列 + 第 3 行第 6 列为
    年级百分位）；学科行只有 grade_percentile（高一表物理同样有百分位
    列，正好透传非 None 值——高二固定列布局物理无百分位列测不到）。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col, text in {1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
                      5: "语文", 7: "物理", 9: "教学班"}.items():
        ws.cell(row=2, column=col, value=text)
    for col, text in {5: "分数", 6: "年级百分位", 7: "分数",
                      8: "年级百分位"}.items():
        ws.cell(row=3, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def test_teaching_grade1_percentile_passthrough(client, v1_seed, tmp_path):
    """teaching 模式旧模板：任教学科（物理）年级百分位透传；非任教学科
    照旧丢弃；名次键恒 None。"""
    row = {
        1: ALIAS_WU_T, 2: 8, 3: 1, 4: "秦甲",
        5: 70, 6: "50%",      # 语文：非任教学科，将被过滤
        7: 78, 8: "45%",      # 物理：分数 + 百分位
        9: "高二8班(教)",
    }
    path = _write_grade1_teaching_xlsx(tmp_path / "FIX01教学明细.xlsx", [row])
    prev = _upload(
        client,
        path,
        "teaching",
        {"teaching_class_id": v1_seed.t8_id, "exam_name": EXAM_TEACHING,
         "exam_date": EXAM_DATE},
    )
    assert prev.status_code == 200, prev.text
    item = prev.json()["items"][0]
    assert item["kind"] == "student_scores" and item["parsed_ok"], item
    assert any("已忽略非任教学科列" in w for w in item["warnings"])
    result = _confirm(client, prev)
    assert result["imported"] == 1, result

    from app.db import workspace_models as wm

    db = _db()
    try:
        facts = (
            db.query(wm.ScoreFact)
            .filter_by(exam_name=EXAM_TEACHING, data_domain="teaching")
            .all()
        )
        assert len(facts) == 1
        fact = facts[0]
        assert fact.identity_id == v1_seed.wu_t_id
        assert fact.subject == "物理" and fact.total_type is None
        assert fact.score == 78.0 and fact.grade_score is None
        assert fact.grade_percentile == 0.45  # "45%" → 0.45
        assert fact.xueji_rank is None and fact.grade_rank is None
    finally:
        db.close()
