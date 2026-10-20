"""P3 名次段表导入 + 段位分布端点 + 均分表班型排名 + 明细扩展字段。

2026-10-19 新增（契约 p3-imports-analysis.md §1/§2.1）：
- 名次段表（学校《六门/三门总分名次段表》）multipart 导入全链路：
  preview 判型/消息 → confirm 落 workspace_rank_band → 再导幂等 skip；
- GET /homeroom/analysis/exams/{exam}/rank-bands：两口径可选、年级行合计、
  无导入时 classes 空列表（不 404、不回退全年级）；
- 班级均分表：新格式指标包（均分/及格率/高分率/低分率/标准差）+ 排名按
  平行班/实验班分组各排（同分同名次）；旧格式裸数值归一化为 {均分}；
- 学生明细端点扩展：grade_scores/percentiles/ranks/xueji/class_num；
- 导入匹配修复回归：legacy 无 scope 学号别名（同学号+同名+本班在册）
  自动接续旧身份，不再新建第二身份（2026-10 重名事故防复发）。

xlsx 样本用 openpyxl 程序化生成，造表 helper 与 test_p3_imports_e01_e02_e05.py
同源复制。
"""

from datetime import date
from pathlib import Path
from urllib.parse import quote

import pytest

API = "/api/v1"

EXAM_RB = "P3RB名次段考"
EXAM_AVG = "P3RB均分考"
EXAM_STU = "P3RB明细考"
EXAM_LEGACY = "P3RB接续考"

_HEADERS = {
    1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
    5: "语文", 6: "数学", 7: "英语",
    8: "物理", 9: "物理等级分",
    20: "+3总分", 21: "主三门总分", 22: "主三门百分位", 23: "主三门学籍排名",
    24: "3+3总分", 25: "3+3百分位", 26: "3+3学籍排名",
}


def _student_row(sid, name, *, class_num=6, xueji=1, chinese=None, math=None,
                 english=None, physics=None, physics_grade=None,
                 plus3=None, main3=None, main3_pct=None, main3_rank=None,
                 total3=None, total3_pct=None, total3_rank=None):
    data = {1: sid, 2: class_num, 3: xueji, 4: name}
    for col, val in ((5, chinese), (6, math), (7, english), (8, physics),
                     (9, physics_grade), (20, plus3), (21, main3), (22, main3_pct),
                     (23, main3_rank), (24, total3), (25, total3_pct), (26, total3_rank)):
        if val is not None:
            data[col] = val
    return data


def _write_detail_xlsx(path: Path, rows) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col, text in _HEADERS.items():
        ws.cell(row=2, column=col, value=text)
    for offset, data in enumerate(rows, start=4):
        for col, val in data.items():
            if val is not None:
                ws.cell(row=offset, column=col, value=val)
    wb.save(str(path))
    return path


def _write_rank_band_xlsx(path: Path, title: str, classes: list[tuple[int, str, list[int]]],
                          bands: list[str]) -> Path:
    """名次段表：row1 合并标题、row2 班级/班主任/段列、row3+ 每班一行。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(bands))
    for col, text in enumerate(["班级", "班主任", *bands], start=1):
        ws.cell(row=2, column=col, value=text)
    for offset, (class_num, teacher, counts) in enumerate(classes, start=3):
        ws.cell(row=offset, column=1, value=class_num)
        ws.cell(row=offset, column=2, value=teacher)
        for idx, count in enumerate(counts, start=3):
            ws.cell(row=offset, column=idx, value=count)
    wb.save(str(path))
    return path


def _write_new_format_averages(path: Path) -> Path:
    """图1 格式均分表：平行班×2 + 实验班×2，总分五子列 + 校名次列。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    tops = ["班级类型", "班级", "班主任", "语文", "数学", "英语",
            "物理", "物理", "化学", "化学",
            "3+3总分", "3+3总分", "3+3总分", "3+3总分", "3+3总分",
            "主三门总分", "主三门总分", "主三门总分", "主三门总分", "主三门总分",
            "3+3校名次", "主三门校名次"]
    subs = ["", "", "", "", "", "",
            "原始分", "等级分", "原始分", "等级分",
            "均分", "及格率", "高分率", "低分率", "标准差",
            "均分", "及格率", "高分率", "低分率", "标准差", "", ""]
    ws.append(tops)
    ws.append(subs)
    # 平行班：6 班均分高、7 班低；实验班：8 班高、9 班低（两类型分别排）
    ws.append(["平行班", 6, "周老师", 90.0, 91.0, 92.0, 70.0, 66.0, 65.0, 61.0,
               440.0, 0.95, 0.60, 0.05, 40.0, 273.0, 96.0, 50.0, 4.0, 30.0, 1, 1])
    ws.append([None, 7, "吴老师", 80.0, 81.0, 82.0, 60.0, 56.0, 55.0, 51.0,
               400.0, 0.85, 0.50, 0.10, 40.1, 243.0, 88.0, 40.0, 11.0, 35.5, 2, 2])
    ws.append(["实验班", 8, "郑老师", 95.0, 96.0, 97.0, 80.0, 76.0, 75.0, 71.0,
               480.0, 0.98, 0.70, 0.02, 30.0, 288.0, 99.0, 60.0, 1.0, 20.0, 1, 1])
    ws.append([None, 9, "王老师", 85.0, 86.0, 87.0, 68.0, 64.0, 63.0, 59.0,
               420.0, 0.90, 0.55, 0.08, 35.0, 258.0, 92.0, 45.0, 8.0, 28.0, 2, 2])
    wb.save(str(path))
    return path


def _upload(client, paths: list[Path], mode: str, extra: dict | None = None):
    data = {"mode": mode}
    if extra:
        data.update({k: str(v) for k, v in extra.items()})
    files = []
    for path in paths:
        files.append((
            "files",
            (path.name, path.read_bytes(),
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ))
    return client.post(f"{API}/imports/preview", data=data, files=files)


def _db():
    from app.db.models import SessionLocal

    return SessionLocal()


def test_rank_band_import_roundtrip_and_endpoint(client, v1_seed, tmp_path):
    """两口径名次段表导入 → 端点读回；重导幂等。"""
    s = v1_seed
    bands = ["1-40", "41-80", "81-120", "121-160"]
    p33 = _write_rank_band_xlsx(
        tmp_path / "六门总分名次段表一（在籍）.xlsx",
        "六门总分名次段表一（在籍）",
        [(5, "周老师", [1, 2, 3, 4]), (6, "王老师", [2, 2, 2, 2]), (8, "郑老师", [3, 1, 0, 0])],
        bands,
    )
    main3 = _write_rank_band_xlsx(
        tmp_path / "三门总分名次段表一（在籍）.xlsx",
        "三门总分名次段表一（在籍）",
        [(5, "周老师", [0, 3, 3, 4]), (6, "王老师", [1, 2, 3, 2]), (8, "郑老师", [2, 2, 1, 0])],
        bands,
    )
    preview = _upload(client, [p33, main3], "homeroom", {
        "class_id": s.h6_id, "exam_name": EXAM_RB, "exam_date": "2025-12-01",
    })
    assert preview.status_code == 200, preview.text
    items = preview.json()["items"]
    assert [it["kind"] for it in items] == ["rank_bands", "rank_bands"]
    assert items[0]["row_count"] == 3 and "3+3" in items[0]["message"]
    assert "主三门" in items[1]["message"] and "23" in items[1]["message"]  # 10+8+5 人

    confirmed = client.post(f"{API}/imports/confirm", json={"token": preview.json()["token"]})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["imported"] == 6  # 3 班 × 2 口径

    # 端点：默认口径=3+3（preferred 顺序），班级按班号升序，年级行=各段求和
    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote(EXAM_RB)}/rank-bands",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_type"] == "3+3"
    assert body["available_total_types"] == ["3+3", "主三门"]
    assert body["current_class_num"] == 6
    assert [c["class_num"] for c in body["classes"]] == [5, 6, 8]
    cls6 = body["classes"][1]
    assert cls6["teacher_name"] == "王老师" and cls6["total"] == 8
    assert [b["count"] for b in cls6["bands"]] == [2, 2, 2, 2]
    assert cls6["bands"][0]["label"] == "1-40" and cls6["bands"][0]["lo"] == 1 and cls6["bands"][0]["hi"] == 40
    assert [b["count"] for b in body["grade_bands"]] == [6, 5, 5, 6]

    # 口径切换
    r2 = client.get(
        f"{API}/homeroom/analysis/exams/{quote(EXAM_RB)}/rank-bands",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id, "total_type": "主三门"},
    )
    assert r2.status_code == 200, r2.json()
    assert r2.json()["total_type"] == "主三门"
    assert [b["count"] for b in r2.json()["grade_bands"]] == [3, 7, 7, 6]

    # 重导幂等：imported=0、全部 skipped
    preview2 = _upload(client, [p33, main3], "homeroom", {
        "class_id": s.h6_id, "exam_name": EXAM_RB, "exam_date": "2025-12-01",
    })
    confirmed2 = client.post(f"{API}/imports/confirm", json={"token": preview2.json()["token"]})
    assert confirmed2.status_code == 200, confirmed2.text
    assert confirmed2.json()["imported"] == 0
    assert confirmed2.json()["skipped"] == 6


def test_rank_bands_empty_state_for_exam_without_import(client, v1_seed):
    """无导入数据的考试：200 + classes 空列表（上传引导空态，不 404）。"""
    from tests.v1.conftest import EXAM_E1

    s = v1_seed
    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote(EXAM_E1)}/rank-bands",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["classes"] == [] and body["grade_bands"] == []
    assert body["available_total_types"] == [] and body["total_type"] is None


def test_class_averages_type_grouped_ranks_and_metric_pack(client, v1_seed, tmp_path):
    """新格式均分表：指标包下发 + 排名按平行班/实验班分别计算。"""
    s = v1_seed
    averages = _write_new_format_averages(tmp_path / "均分表.xlsx")
    preview = _upload(client, [averages], "homeroom", {
        "class_id": s.h6_id, "exam_name": EXAM_AVG, "exam_date": "2025-12-02",
    })
    assert preview.status_code == 200, preview.text
    assert [it["kind"] for it in preview.json()["items"]] == ["class_averages"]
    confirmed = client.post(f"{API}/imports/confirm", json={"token": preview.json()["token"]})
    assert confirmed.status_code == 200, confirmed.text

    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote(EXAM_AVG)}/class-averages",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    by_num = {row["class_num"]: row for row in body["rows"]}
    assert set(by_num) == {6, 7, 8, 9}

    # 指标包：及格率 0.95 归一化为 95.0；校名次（文件值）不下发
    entry = by_num[6]["totals"]["3+3"]
    assert entry["均分"] == 440.0 and entry["及格率"] == 95.0
    assert entry["高分率"] == 60.0 and entry["低分率"] == 5.0 and entry["标准差"] == 40.0
    assert "校名次" not in entry

    # 班型分组排名：平行班内 6 班第 1、7 班第 2；实验班内 8 班第 1、9 班第 2
    # （全校口径下 7 班第 4，但班型内只与平行班比较）
    assert by_num[6]["total_ranks"]["3+3"] == 1 and by_num[7]["total_ranks"]["3+3"] == 2
    assert by_num[8]["total_ranks"]["3+3"] == 1 and by_num[9]["total_ranks"]["3+3"] == 2
    assert by_num[6]["total_ranks"]["主三门"] == 1 and by_num[8]["total_ranks"]["主三门"] == 1

    # 选科原始/等级子列照常下发
    assert by_num[6]["subjects"]["物理_原始"] == 70.0 and by_num[6]["subjects"]["物理_等级"] == 66.0


def test_class_averages_old_numeric_totals_normalized(client, v1_seed):
    """旧格式裸数值（ORM 直种）在端点归一化为 {均分: v}，排名照常分组。"""
    from app.db import workspace_models as wm

    s = v1_seed
    db = _db()
    for num, class_type, avg in ((1, "平行班", 300.0), (2, "平行班", 310.0), (3, "实验班", 280.0)):
        db.add(wm.WorkspaceClassAverage(
            data_domain="homeroom", academic_year_id=s.ay_id,
            exam_name="P3RB旧均分考", exam_date=date(2025, 12, 3),
            grade=2, class_type=class_type, class_num=num,
            teacher_name="旧老师",
            subject_averages={}, total_averages={"主三门": avg},
            source="p3-rank-bands-test",
        ))
    db.commit()
    db.close()

    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote('P3RB旧均分考')}/class-averages",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id},
    )
    assert r.status_code == 200, r.text
    rows = {row["class_num"]: row for row in r.json()["rows"]}
    assert rows[2]["totals"]["主三门"] == {"均分": 310.0}
    assert rows[1]["total_ranks"]["主三门"] == 2 and rows[2]["total_ranks"]["主三门"] == 1
    assert rows[3]["total_ranks"]["主三门"] == 1  # 实验班单独排


def test_students_endpoint_extended_fields(client, v1_seed, tmp_path):
    """明细表扩展字段：等级分/百分位/校名次/学籍/班级号。"""
    from tests.v1.conftest import ALIAS_JIA_H

    s = v1_seed
    detail = _write_detail_xlsx(tmp_path / "明细.xlsx", [
        _student_row(ALIAS_JIA_H, "秦甲", chinese=88, math=92, english=95,
                     physics=90, physics_grade=86, plus3=176, main3=275,
                     main3_pct=91.2, main3_rank=12, total3=451, total3_pct=88.5,
                     total3_rank=15),
    ])
    preview = _upload(client, [detail], "homeroom", {
        "class_id": s.h6_id, "exam_name": EXAM_STU, "exam_date": "2025-12-04",
    })
    assert preview.status_code == 200, preview.text
    confirmed = client.post(f"{API}/imports/confirm", json={"token": preview.json()["token"]})
    assert confirmed.status_code == 200, confirmed.text

    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote(EXAM_STU)}/students",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id},
    )
    assert r.status_code == 200, r.text
    students = r.json()["students"]
    jia = next(st for st in students if st["alias"] == ALIAS_JIA_H)
    assert jia["scores"]["物理"] == 90.0
    assert jia["grade_scores"]["物理"] == 86.0
    # 百分位/校名次挂口径键；xueji/class_num 为身份级字段
    assert jia["percentiles"]["主三门"] == pytest.approx(0.912)
    assert jia["ranks"]["3+3"] == 15 and jia["ranks"]["主三门"] == 12
    assert jia["xueji"] == 1 and jia["class_num"] == 6


def test_same_class_legacy_alias_auto_continues(client, v1_seed, tmp_path):
    """重名事故防复发：无 scope 的 legacy 学号别名 + 同名 + 本班在册 →
    导入接续旧身份，不新建第二身份。"""
    from app.db import workspace_models as wm

    s = v1_seed
    alias_value = "7299901"
    db = _db()
    legacy = wm.WsStudentIdentity(
        data_domain="homeroom", display_name="秦传", note="migration:h", xueji=3,
    )
    db.add(legacy)
    db.flush()
    db.add(wm.WsStudentAlias(
        identity_id=legacy.id, alias_value=alias_value, data_domain="homeroom",
        source="migration:h", academic_year_id=None, alias_scope="none",
    ))
    db.add(wm.Enrollment(
        admin_class_id=s.h6_id, identity_id=legacy.id, status="active",
        valid_from=date(2025, 9, 1),
    ))
    db.commit()
    legacy_id = legacy.id
    db.close()

    detail = _write_detail_xlsx(tmp_path / "接续明细.xlsx", [
        _student_row(alias_value, "秦传", chinese=70, math=71, english=72, main3=213),
    ])
    preview = _upload(client, [detail], "homeroom", {
        "class_id": s.h6_id, "exam_name": EXAM_LEGACY, "exam_date": "2025-12-05",
    })
    assert preview.status_code == 200, preview.text
    item = preview.json()["items"][0]
    # 三重证据齐备 → 直接 known 接续，不进候选不新建
    assert item["known_students"] == 1 and item["new_students"] == []
    assert item["identity_candidates"] == []

    confirmed = client.post(f"{API}/imports/confirm", json={"token": preview.json()["token"]})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["students_created"] == 0

    # 成绩挂在旧身份上；库里没有第二个同名身份
    db = _db()
    facts = db.query(wm.ScoreFact).filter_by(
        academic_year_id=s.ay_id, exam_name=EXAM_LEGACY, identity_id=legacy_id,
    ).all()
    assert len(facts) == 4  # 语数英 + 主三门
    dup = db.query(wm.WsStudentIdentity).filter_by(
        data_domain="homeroom", display_name="秦传",
    ).all()
    assert len(dup) == 1
    db.close()

    # 学籍代码按本次文件值覆盖更新（2026-10-20 裁决：文件有值即写；
    # 直种原值 3，文件第 3 列 xueji=1 → 更新为 1）
    db = _db()
    assert db.get(wm.WsStudentIdentity, legacy_id).xueji == 1
    db.close()

    # 学生表一行、有成绩（不再出现「旧身份全—+新身份重名」）
    r = client.get(
        f"{API}/homeroom/analysis/exams/{quote(EXAM_LEGACY)}/students",
        params={"academic_year_id": s.ay_id, "class_id": s.h6_id},
    )
    assert r.status_code == 200, r.text
    students = r.json()["students"]
    chuan = [st for st in students if st["name"] == "秦传"]
    assert len(chuan) == 1 and chuan[0]["totals"]["主三门"] == 213.0
