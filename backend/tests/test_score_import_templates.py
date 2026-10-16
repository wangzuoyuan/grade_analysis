"""Fill the shipped downloads and read them with the actual import adapters."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.ingest.teaching_platform import parse
from app.ingest.ws_parse import parse_homeroom_file, parse_teaching_file


TEMPLATES = Path(__file__).resolve().parents[2] / "frontend/public/templates/scores/v1"
DETAILS = ["teaching-g1-students", "teaching-g23-students", "homeroom-g1-students", "homeroom-g23-students"]
QUALITY = ["teaching-g1-quality", "teaching-g23-teaching-quality", "teaching-g23-admin-quality"]
AVERAGES = ["homeroom-g1-averages", "homeroom-g23-averages"]


@pytest.mark.parametrize("template", DETAILS + QUALITY + AVERAGES)
def test_downloads_are_blank_single_sheet_workbooks(template):
    wb = load_workbook(TEMPLATES / f"{template}.xlsx")
    assert len(wb.worksheets) == 1
    ws = wb.active
    header_rows = 3 if template.startswith("homeroom") and template.endswith("students") else 1 if template.startswith("teaching") and template.endswith("students") else 2
    assert all(cell.value is None for row in ws.iter_rows(min_row=header_rows + 1) for cell in row)
    if template.endswith("students"):
        assert ws.cell(header_rows + 1, 1).number_format == "@"


@pytest.mark.parametrize("grade", [1, 2, 3])
def test_platform_detail_download_supports_blank_supplement_and_zero(tmp_path, grade):
    name = f"teaching-g{'1' if grade == 1 else '23'}-students"
    wb = load_workbook(TEMPLATES / f"{name}.xlsx")
    ws = wb.active
    for c, value in enumerate(["000TEST01", "模板测试甲", 0], 1):
        ws.cell(2, c, value)
    target = tmp_path / "details.xlsx"
    wb.save(target)
    row = parse(target.read_bytes(), target.name, grade, "数学")["rows"][0]
    assert row["alias"] == "000TEST01" and row["score"] == 0
    assert row.get("grade_score") is None and row.get("xueji_rank") is None
    assert row.get("grade_percentile") is None
    ws.cell(2, 4, 9 if grade == 1 else 64)
    if grade == 1:
        ws.cell(2, 5, 0.25)
    wb.save(target)
    row = parse(target.read_bytes(), target.name, grade, "数学")["rows"][0]
    assert row["score"] == 0
    if grade == 1:
        assert row["xueji_rank"] == 9 and row["grade_percentile"] == 0.25
    else:
        assert row["grade_score"] == 64


@pytest.mark.parametrize("grade", [1, 2, 3])
@pytest.mark.parametrize("administrative", [False, True])
def test_platform_quality_downloads_preserve_case_and_teacher_rules(tmp_path, grade, administrative):
    name = "teaching-g1-quality" if grade == 1 else f"teaching-g23-{'admin' if administrative else 'teaching'}-quality"
    wb = load_workbook(TEMPLATES / f"{name}.xlsx")
    ws = wb.active
    if grade == 1:
        values = ["模板班型", "TemplateAa", None, 2, 2, 0, 0, 0, 0, 0] + [0] * 8
    else:
        bands = [0] * 22
        bands[6:8] = [2, 1]  # B: both students, 100%.
        values = (["模板班型", "TemplateAa", "模板任课教师"] if administrative else ["TemplateAa", None]) + [2, 2, 0, 61, 0] + bands
    for c, value in enumerate(values, 1):
        ws.cell(3, c, value)
        ws.cell(4, c, value)
    label_col = 2 if grade == 1 or administrative else 1
    teacher_col = label_col + 1
    ws.cell(4, label_col, "TemplateaA")
    ws.cell(4, teacher_col).value = None
    target = tmp_path / "quality.xlsx"
    wb.save(target)
    result = parse(target.read_bytes(), target.name, grade, "数学")
    expected = ["TemplateAa"] if grade != 1 and administrative else ["TemplateAa", "TemplateaA"]
    assert [c["class_label"] for c in result["classes"]] == expected
    assert result["classes"][0]["metrics"]["raw_avg"] == 0
    assert result["classes"][0]["metrics"]["stddev"] == 0
    if grade != 1:
        assert result["classes"][0]["metrics"]["grade_avg"] == 61
        b = next(b for b in result["classes"][0]["metrics"]["bands"] if b["label"] == "B")
        assert b == {"label": "B", "count": 2, "ratio": 1}


@pytest.mark.parametrize("grade", [1, 2, 3])
def test_full_subject_detail_downloads_work_in_both_legacy_adapters(tmp_path, grade):
    source = TEMPLATES / f"homeroom-g{'1' if grade == 1 else '23'}-students.xlsx"
    wb = load_workbook(source)
    ws = wb.active
    values = {1: "000TEST01", 2: 6, 3: 1, 4: "模板测试甲", 5: 0}
    values.update({6: 0.25, 7: 88, 23: 240, 24: 0.5, 25: 9} if grade == 1 else {6: 88, 8: 0, 9: 61, 21: 240, 22: 0.5, 23: 9})
    for c, value in values.items():
        ws.cell(4, c, value)
    # The legacy adapter also uses the grade-bearing download filename.
    filename = f"高{'一' if grade == 1 else '二' if grade == 2 else '三'}全科成绩明细模板.xlsx"
    target = tmp_path / filename
    wb.save(target)
    homeroom = parse_homeroom_file(target, filename, grade, 6, "模板合成考试", "2030-10-10")
    assert homeroom.parsed_ok and homeroom.kind == "student_scores"
    chinese = next(r for r in homeroom.rows if r.get("subject") == "语文")
    assert chinese["score"] == 0 and chinese["alias"] == "000TEST01"
    total = next(r for r in homeroom.rows if r.get("total_type") == "主三门")
    assert total["score"] == 240
    teaching = parse_teaching_file(target, filename, "数学", None, "模板合成考试", "2030-10-10")
    assert teaching.parsed_ok and teaching.rows
    assert all(r.get("subject") == "数学" for r in teaching.rows)
    assert teaching.rows[0]["score"] == 88


@pytest.mark.parametrize("grade", [1, 2, 3])
def test_homeroom_class_average_downloads(tmp_path, grade):
    wb = load_workbook(TEMPLATES / f"homeroom-g{'1' if grade == 1 else '23'}-averages.xlsx")
    ws = wb.active
    for c, value in {1: "模板班型", 2: 6, 3: "模板教师", 4: 0, 5: 88, 13 if grade == 1 else 20: 240}.items():
        ws.cell(3, c, value)
    target = tmp_path / "averages.xlsx"
    wb.save(target)
    result = parse_homeroom_file(target, target.name, grade, 6, "模板合成考试", "2030-10-10")
    assert result.parsed_ok and result.kind == "class_averages"
    row = result.class_averages[0]
    assert row["class_num"] == 6
    assert row["subject_averages"]["语文"] == 0 and row["subject_averages"]["数学"] == 88
    assert row["total_averages"]["主三门"] == 240
