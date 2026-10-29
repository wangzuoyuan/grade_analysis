from pathlib import Path

import pytest
from openpyxl import Workbook

from app.ingest.excel_parser import parse_excel_grade23
from app.ingest.excel_parser import parse_excel_grade1_student_scores
from app.ingest.excel_parser import is_grade23_rank_band_sheet, is_grade23_student_sheet


def make_grade23_student_workbook(path: Path):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "学生成绩明细"
    ws.cell(2, 1, "学号")
    ws.cell(2, 2, "班级")
    ws.cell(2, 3, "学籍")
    ws.cell(2, 4, "姓名")
    ws.cell(2, 20, "+3总分")
    ws.cell(2, 21, "主三门")
    ws.cell(2, 24, "3+3总分")

    # 高二/高三学生成绩表是固定列位解析：5-7 为语数英原始分，
    # 8-19 为选考科目原始/等级分，20-26 为总分/百分位/排名，27-29 为语数英百分位。
    ws.append([])
    ws.append([
        "7240101", "01", "1", "卞幻",
        97, 108, 120,
        48, 52,
        50, 55,
        None, None,
        None, None,
        None, None,
        None, None,
        174,
        325.5, "25.08%", 283,
        499.5, "30.01%", 291,
        "75.08%", "40.12%", "35.45%",
    ])
    ws.append([
        "7240102", "01", "1", "学生乙",
        90, 100, 110,
        None, None,
        45, 50,
        None, None,
        None, None,
        None, None,
        None, None,
        95,
        300, "40.00%", 350,
        395, "45.00%", 420,
        "60.00%", "55.00%", "50.00%",
    ])
    workbook.save(path)


def test_parse_grade23_sample_totals(tmp_path):
    sample = tmp_path / "高二2025学年第二学期期中考试学生成绩明细表.xlsx"
    make_grade23_student_workbook(sample)
    result = parse_excel_grade23(str(sample), grade=2)

    assert result["kind"] == "student_scores"
    assert len(result["students"]) == 2

    totals = {
        row["total_type"]: row
        for row in result["total_scores"]
        if row["student_id"] == "7240101"
    }
    assert totals["主三门"]["total_score"] == 325.5
    assert totals["+3"]["total_score"] == 174
    assert totals["3+3"]["total_score"] == 499.5
    assert totals["3+3"]["xueji_rank"] == 291


def test_parse_grade23_sample_subjects(tmp_path):
    sample = tmp_path / "高二2025学年第二学期期中考试学生成绩明细表.xlsx"
    make_grade23_student_workbook(sample)
    result = parse_excel_grade23(str(sample), grade=2)

    subjects = {
        row["subject"]: row
        for row in result["subject_scores"]
        if row["student_id"] == "7240101"
    }
    assert subjects["语文"]["raw_score"] == 97
    assert subjects["语文"]["grade_percentile"] == 0.7508
    assert subjects["物理"]["raw_score"] == 48
    assert subjects["物理"]["grade_score"] == 52
    assert "生物" not in subjects


# ── 2026-10-24 新版明细表：主三门/3+3 各带「学籍排名+年级排名」两列名次 ──
# 学科列恒在 5–19；总分块扩为 4 列（21-24 / 25-28），语数英百分位移至 29-31。


def _grade23_new_format_rows(ws):
    """两名学生按列号写入（新版布局）。"""
    rows = {
        "7240101": ("06", 1, "卞幻", {5: 100, 6: 110, 7: 120, 8: 70, 9: 80,
                                      20: 174, 21: 325.5, 22: "25.08%", 23: 283, 24: 290,
                                      25: 499.5, 26: "30.01%", 27: 291, 28: 298,
                                      29: "75.08%", 30: "40.12%", 31: "35.45%"}),
        "7240102": ("06", 2, "学生乙", {5: 90, 6: 100, 7: 110, 8: None, 9: None,
                                        20: 95, 21: 300, 22: "40.00%", 23: 350, 24: 358,
                                        25: 395, 26: "45.00%", 27: 420, 28: 428,
                                        29: "60.00%", 30: "55.00%", 31: "50.00%"}),
    }
    for r, (sid, (cls, xueji, name, cols)) in enumerate(rows.items(), start=4):
        ws.cell(r, 1, sid)
        ws.cell(r, 2, cls)
        ws.cell(r, 3, xueji)
        ws.cell(r, 4, name)
        for c, v in cols.items():
            if v is not None:
                ws.cell(r, c, v)


def make_grade23_new_format_two_row(path: Path):
    """新版·两行表头（总标签合并 + 子标签，同高一风格）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生成绩明细"
    for c, t in {1: "学号", 2: "班级", 3: "学籍", 4: "姓名", 5: "语文", 6: "数学", 7: "英语",
                 8: "物理", 10: "化学", 20: "+3总分", 21: "主三门", 25: "3+3总分",
                 29: "语文", 30: "数学", 31: "英语"}.items():
        ws.cell(2, c, t)
    for c, s in {8: "原始分", 9: "等级分", 10: "原始分", 11: "等级分",
                 21: "总分", 22: "年级百分位", 23: "学籍排名", 24: "年级排名",
                 25: "总分", 26: "年级百分位", 27: "学籍排名", 28: "年级排名",
                 29: "年级百分位", 30: "年级百分位", 31: "年级百分位"}.items():
        ws.cell(3, c, s)
    for c in (8, 10):
        ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c + 1)
    ws.merge_cells(start_row=2, start_column=21, end_row=2, end_column=24)
    ws.merge_cells(start_row=2, start_column=25, end_row=2, end_column=28)
    _grade23_new_format_rows(ws)
    wb.save(path)


def make_grade23_new_format_one_row(path: Path):
    """新版·单行表头（整列名：主三门总分/主三门学籍排名/主三门年级排名…）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生成绩明细"
    for c, t in {1: "学号", 2: "班级", 3: "学籍", 4: "姓名", 5: "语文", 6: "数学", 7: "英语",
                 8: "物理", 9: "等级分"}.items():
        ws.cell(2, c, t)  # 5-19 学科列本测试只填到 9；其余按需
    for c, t in {20: "+3总分", 21: "主三门总分", 22: "主三门年级百分位", 23: "主三门学籍排名",
                 24: "主三门年级排名", 25: "3+3总分", 26: "3+3年级百分位", 27: "3+3学籍排名",
                 28: "3+3年级排名", 29: "语文年级百分位", 30: "数学年级百分位", 31: "英语年级百分位"}.items():
        ws.cell(2, c, t)
    _grade23_new_format_rows(ws)
    wb.save(path)


@pytest.mark.parametrize("builder", [
    make_grade23_new_format_two_row,
    make_grade23_new_format_one_row,
])
def test_parse_grade23_new_format_dual_ranks(tmp_path, builder):
    """新版明细表（两行/单行表头两种形态）双列名次全解析；
    学籍列照进 students。"""
    sample = tmp_path / "高二2026学年第一学期10月考学生成绩明细表.xlsx"
    builder(sample)
    assert is_grade23_student_sheet(sample)

    result = parse_excel_grade23(str(sample), grade=2)
    assert result["kind"] == "student_scores"

    totals = {row["total_type"]: row for row in result["total_scores"] if row["student_id"] == "7240101"}
    assert totals["主三门"]["total_score"] == 325.5
    assert abs(totals["主三门"]["grade_percentile"] - 0.2508) < 1e-9
    assert totals["主三门"]["xueji_rank"] == 283
    assert totals["主三门"]["grade_rank"] == 290
    assert totals["3+3"]["total_score"] == 499.5
    assert totals["3+3"]["xueji_rank"] == 291
    assert totals["3+3"]["grade_rank"] == 298
    assert totals["+3"]["total_score"] == 174

    subjects = {row["subject"]: row for row in result["subject_scores"] if row["student_id"] == "7240101"}
    assert abs(subjects["语文"]["grade_percentile"] - 0.7508) < 1e-9
    assert abs(subjects["数学"]["grade_percentile"] - 0.4012) < 1e-9
    assert abs(subjects["英语"]["grade_percentile"] - 0.3545) < 1e-9
    assert subjects["物理"]["raw_score"] == 70
    assert subjects["物理"]["grade_score"] == 80

    xueji = {s["student_id"]: s["xueji"] for s in result["students"]}
    assert xueji == {"7240101": 1, "7240102": 2}


def test_parse_grade23_new_format_partial_columns(tmp_path):
    """极简新格式文件（只有主三门块 + 语义化列名，无 3+3 块）：
    识别到的槽照解析、缺失块不兜底不报错。"""
    wb = Workbook()
    ws = wb.active
    for c, t in {1: "学号", 2: "班级", 3: "学籍", 4: "姓名",
                 21: "主三门总分", 22: "主三门年级百分位", 23: "主三门学籍排名",
                 24: "主三门年级排名"}.items():
        ws.cell(2, c, t)
    ws.cell(4, 1, "7240101")
    ws.cell(4, 2, "06")
    ws.cell(4, 3, 1)
    ws.cell(4, 4, "卞幻")
    ws.cell(4, 21, 325.5)
    ws.cell(4, 23, 283)
    ws.cell(4, 24, 290)
    sample = tmp_path / "高二极简新版明细.xlsx"
    wb.save(sample)
    assert is_grade23_student_sheet(sample)

    result = parse_excel_grade23(str(sample), grade=2)
    totals = {row["total_type"]: row for row in result["total_scores"]}
    assert totals["主三门"]["xueji_rank"] == 283
    assert totals["主三门"]["grade_rank"] == 290
    assert "+3" not in totals
    assert "3+3" not in totals


def make_grade23_20261028_variant(path: Path):
    """2026-10-28 学校再版明细表（真实文件复刻，匿名化数据）：
    主三门合并头标「三门」、3+3 块标「所有（3+3）」、子标签写
    「年级百分比」「百分比」（非「年级百分位」）。修复前三处错位：
    主三门整块丢失 / 语数英百分位落到旧固定列 27-29 装进名次 /
    3+3 百分位兜底撞总分列（score÷100 假百分位）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生成绩明细"
    for c, t in {1: "学号", 2: "班级", 3: "学籍", 4: "姓名", 5: "语文", 6: "数学", 7: "英语",
                 8: "物理", 10: "化学", 12: "生物", 14: "政治", 16: "历史", 18: "地理",
                 20: "+3总分", 21: "三门", 25: "所有（3+3）",
                 29: "语文", 30: "数学", 31: "英语"}.items():
        ws.cell(2, c, t)
    for c, s in {8: "原始", 9: "等级", 10: "原始", 11: "等级", 12: "原始", 13: "等级",
                 14: "原始", 15: "等级", 16: "原始", 17: "等级", 18: "原始", 19: "等级",
                 21: "总分", 22: "年级百分比", 23: "学籍排名", 24: "年级排名",
                 25: "总分", 26: "年级百分比", 27: "学籍排名", 28: "年级排名",
                 29: "百分比", 30: "百分比", 31: "百分比"}.items():
        ws.cell(3, c, s)
    ws.merge_cells(start_row=2, start_column=21, end_row=2, end_column=24)
    ws.merge_cells(start_row=2, start_column=25, end_row=2, end_column=28)
    # 两行数据：数值取真实文件形状（名次 3 位数、百分位 0-1 小数）
    for r, (sid, cls, xj, name, cols) in {
        4: ("7240101", "06", 1, "卞幻", {5: 89, 6: 80, 7: 88, 8: 58, 9: 52, 10: 41, 11: 46,
                                         20: 165, 21: 257, 22: 0.7781, 23: 477, 24: 480,
                                         25: 422, 26: 0.8206, 27: 503, 28: 505,
                                         29: 0.6476, 30: 0.7798, 31: 0.6977}),
        5: ("7240102", "06", 2, "学生乙", {5: 86, 6: 76, 7: 77,
                                          20: 159, 21: 239, 22: 0.7863, 23: 485, 24: 488,
                                          25: 415, 26: 0.8288, 27: 522, 28: 524,
                                          29: 0.7749, 30: 0.6234, 31: 0.5566}),
    }.items():
        ws.cell(r, 1, sid)
        ws.cell(r, 2, cls)
        ws.cell(r, 3, xj)
        ws.cell(r, 4, name)
        for c, v in cols.items():
            ws.cell(r, c, v)
    wb.save(path)


def test_parse_grade23_20261028_variant_labels(tmp_path):
    """2026-10-28 变体标签全解析：三门/所有（3+3）/百分比 词汇兼容，
    主三门块在场、百分位各归各列（修复前主三门丢失+百分位装名次/总分）。"""
    sample = tmp_path / "高二再版明细表.xlsx"
    make_grade23_20261028_variant(sample)
    assert is_grade23_student_sheet(sample)

    result = parse_excel_grade23(str(sample), grade=2)
    assert result["kind"] == "student_scores"

    totals = {row["total_type"]: row for row in result["total_scores"] if row["student_id"] == "7240101"}
    assert set(totals) == {"+3", "主三门", "3+3"}, "主三门整块不得再丢失"
    assert totals["主三门"]["total_score"] == 257
    assert abs(totals["主三门"]["grade_percentile"] - 0.7781) < 1e-9
    assert totals["主三门"]["xueji_rank"] == 477
    assert totals["主三门"]["grade_rank"] == 480
    assert totals["3+3"]["total_score"] == 422
    assert abs(totals["3+3"]["grade_percentile"] - 0.8206) < 1e-9, "3+3 百分位不得再撞总分列（422÷100=4.22）"
    assert totals["3+3"]["xueji_rank"] == 503
    assert totals["3+3"]["grade_rank"] == 505

    subjects = {row["subject"]: row for row in result["subject_scores"] if row["student_id"] == "7240101"}
    assert abs(subjects["语文"]["grade_percentile"] - 0.6476) < 1e-9, "语文百分位不得再装进 3+3 学籍排名（503÷100）"
    assert abs(subjects["数学"]["grade_percentile"] - 0.7798) < 1e-9, "数学百分位不得再装进 3+3 年级排名（505÷100）"
    assert abs(subjects["英语"]["grade_percentile"] - 0.6977) < 1e-9, "英语百分位不得再装进语文百分位列"
    assert subjects["物理"]["grade_score"] == 52


def test_parse_grade1_percentiles(tmp_path):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "学生成绩明细"
    ws.append(["学生成绩（在籍）"])
    ws.append(["学号", "班级", "学籍", "姓名", "语文", None, "数学", None, "主三门", None, None, None])
    ws.append([None, None, None, None, "分数", "年级百分位", "分数", "年级百分位", "总分", "年级百分位", "学籍排名", "年级排名"])
    ws.merge_cells("A2:A3")
    ws.merge_cells("B2:B3")
    ws.merge_cells("C2:C3")
    ws.merge_cells("D2:D3")
    ws.merge_cells("E2:F2")
    ws.merge_cells("G2:H2")
    ws.merge_cells("I2:L2")
    ws.append(["7240101", "01", "1", "卞幻", 92, "45.01%", 106, "66.84%", 306, "47.88%", 283, 283])
    path = tmp_path / "2024级2024学年第二学期期中考试.xlsx"
    workbook.save(path)

    result = parse_excel_grade1_student_scores(path)

    subjects = {
        row["subject"]: row
        for row in result["subject_scores"]
        if row["student_id"] == "7240101"
    }
    totals = {
        row["total_type"]: row
        for row in result["total_scores"]
        if row["student_id"] == "7240101"
    }
    assert subjects["语文"]["grade_percentile"] == 0.4501
    assert subjects["数学"]["grade_percentile"] == 0.6684
    assert totals["主三门"]["grade_percentile"] == 0.4788
    assert totals["主三门"]["xueji_rank"] == 283
    assert totals["主三门"]["grade_rank"] == 283


def test_parse_grade23_class_average_workbook(tmp_path):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "班级均分"
    ws.append(
        [
            "班型",
            "班级",
            "班主任",
            "语文",
            "数学",
            "英语",
            "物理",
            None,
            "化学",
            None,
            "加3同均分",
            "主三门",
            "3+3总分",
        ]
    )
    ws.append([None, None, None, None, None, None, "原始", "等级", "原始", "等级", None, None, None])
    ws.append(["平行班", "01", "张老师", 101.2, 108.3, 110.4, 55.1, 61.2, 58.3, 62.4, 180.5, 319.9, 500.4])
    ws.append([None, "02", "李老师", 99.1, 103.2, 109.3, 53.1, 59.2, 57.3, 60.4, 176.5, 311.6, 488.1])
    path = tmp_path / "高二2025学年第二学期期中考试班级均分表.xlsx"
    workbook.save(path)

    result = parse_excel_grade23(str(path), grade=2)

    assert result["kind"] == "class_averages"
    assert len(result["class_averages"]) == 2
    first = result["class_averages"][0]
    assert first["class_type"] == "平行班"
    assert first["class_num"] == 1
    assert first["teacher_name"] == "张老师"
    assert first["subject_averages"]["语文"] == 101.2
    assert first["subject_averages"]["物理_等级"] == 61.2
    # 总分口径升级为「口径→指标包」dict：旧格式无子列时均分落 {均分: v}
    assert first["total_averages"]["+3"] == {"均分": 180.5}
    assert first["total_averages"]["主三门"] == {"均分": 319.9}
    assert first["total_averages"]["3+3"] == {"均分": 500.4}


def make_grade23_new_format_average_workbook(path: Path):
    # 新格式学校均分表：总分列挂 均分/及格率/高分率/低分率/标准差 子列 + 校名次顶级列
    workbook = Workbook()
    ws = workbook.active
    ws.title = "班级均分"
    ws.append(
        [
            "班级类型", "班级", "班主任", "语文", "数学", "英语", "物理", None,
            "3+3总分", None, None, None, None,
            "主三门总分", None, None, None, None,
            "3+3校名次", "主三门校名次",
        ]
    )
    ws.append(
        [None, None, None, None, None, None, "原始", "等级"]
        + ["均分", "及格率", "高分率", "低分率", "标准差"]
        + ["均分", "及格率", "高分率", "低分率", "标准差"]
        + [None, None]
    )
    ws.append(
        ["平行班", "01", "张老师", 101.2, 108.3, 110.4, 55.1, 61.2,
         499.5, 0.91, 0.12, 0.03, 45.67, 325.5, 95, 0.2, 0, 38.5, 3, 5]
    )
    ws.append(
        [None, "02", "李老师", 99.1, 103.2, 109.3, 53.1, 59.2,
         488.1, 0.85, 0.1, 0.02, 44.2, 311.6, 0.93, 0.15, 0.05, 37.9, 8, 6]
    )
    ws.merge_cells("A3:A4")  # 班级类型纵向合并：第二行靠继承
    workbook.save(path)


def test_parse_grade23_new_format_class_average_workbook(tmp_path):
    sample = tmp_path / "高二合成名次段期初考班级均分表.xlsx"
    make_grade23_new_format_average_workbook(sample)
    result = parse_excel_grade23(str(sample), grade=2)

    assert result["kind"] == "class_averages"
    assert len(result["class_averages"]) == 2
    first, second = result["class_averages"]
    # 选考科目双列：物理_原始/物理_等级 分列入包
    assert first["subject_averages"]["物理_原始"] == 55.1
    assert first["subject_averages"]["物理_等级"] == 61.2
    # 3+3 完整指标包：及格率 0.91（比例）自动 ×100 存百分数；校名次顶级列入包
    assert first["total_averages"]["3+3"] == {
        "均分": 499.5, "及格率": 91.0, "高分率": 12.0, "低分率": 3.0,
        "标准差": 45.67, "校名次": 3,
    }
    # 主三门包：及格率 95 已是百分数原样保留；校名次同口径入包
    assert first["total_averages"]["主三门"]["均分"] == 325.5
    assert first["total_averages"]["主三门"]["及格率"] == 95.0
    assert first["total_averages"]["主三门"]["校名次"] == 5
    # 班级类型合并单元格：第二行空值继承「平行班」
    assert second["class_type"] == "平行班"
    assert second["total_averages"]["3+3"]["及格率"] == 85.0
    assert second["total_averages"]["主三门"]["校名次"] == 6


def make_grade23_rank_band_workbook(path: Path, title: str):
    # 名次段表：第 1 行合并标题，第 2 行 班级/班主任/段列，第 3 行起每班一行
    workbook = Workbook()
    ws = workbook.active
    ws.title = "名次段"
    ws.append([title])
    ws.append(["班级", "班主任", "1-40", "41-80", "81-120", "121-160"])
    ws.append(["01", "张老师", 5, 8, 12, 3])
    ws.append(["02", "李老师", 4, 9, 10, 6])
    ws.append(["合计", "—", 9, 17, 22, 9])  # 非数字班级行：应被忽略
    workbook.save(path)


def test_parse_grade23_rank_band_workbook(tmp_path):
    sample = tmp_path / "高二合成名次段期初考名次段.xlsx"
    make_grade23_rank_band_workbook(sample, "六门总分名次段表一（在籍）")
    result = parse_excel_grade23(str(sample), grade=2)

    # 标题含「六门」→ 3+3 口径；合计行非数字班号被忽略，只留两个班
    assert result["kind"] == "rank_bands"
    band = result["rank_band"]
    assert band["total_type"] == "3+3"
    assert band["title"] == "六门总分名次段表一（在籍）"
    assert len(band["classes"]) == 2
    first = band["classes"][0]
    assert first["class_num"] == 1
    assert first["teacher_name"] == "张老师"
    assert first["bands"] == [
        {"label": "1-40", "lo": 1, "hi": 40, "count": 5},
        {"label": "41-80", "lo": 41, "hi": 80, "count": 8},
        {"label": "81-120", "lo": 81, "hi": 120, "count": 12},
        {"label": "121-160", "lo": 121, "hi": 160, "count": 3},
    ]

    # 同构文件标题含「三门」→ 主三门口径
    sample_main3 = tmp_path / "高二合成名次段期初考名次段2.xlsx"
    make_grade23_rank_band_workbook(sample_main3, "三门总分名次段表一（在籍）")
    result_main3 = parse_excel_grade23(str(sample_main3), grade=3)
    assert result_main3["kind"] == "rank_bands"
    assert result_main3["rank_band"]["total_type"] == "主三门"


def test_grade23_rank_band_detection_is_exclusive(tmp_path):
    band = tmp_path / "高二合成名次段期初考名次段.xlsx"
    make_grade23_rank_band_workbook(band, "六门总分名次段表一（在籍）")
    workbook = Workbook()
    ws = workbook.active
    ws.append(["班型", "班级", "班主任", "语文", "数学", "英语", "物理", None, "加3同均分", "主三门", "3+3总分"])
    ws.append([None, None, None, None, None, None, "原始", "等级", None, None, None])
    ws.append(["平行班", "01", "张老师", 101.2, 108.3, 110.4, 55.1, 61.2, 180.5, 319.9, 500.4])
    averages = tmp_path / "高二2025学年第二学期期中考试班级均分表.xlsx"
    workbook.save(averages)

    # 判型互斥：名次段表不误判为明细表，均分表不误判为名次段表，分流各归其位
    assert is_grade23_student_sheet(band) is False
    assert is_grade23_rank_band_sheet(band) is True
    assert is_grade23_rank_band_sheet(averages) is False
    assert parse_excel_grade23(str(band), grade=2)["kind"] == "rank_bands"
    assert parse_excel_grade23(str(averages), grade=2)["kind"] == "class_averages"
