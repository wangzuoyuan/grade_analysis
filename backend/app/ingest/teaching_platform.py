"""教学平台单科文件。纯解析：不访问数据库，不用文件名猜考试/年级。"""
from io import BytesIO
import math
import re
from openpyxl import load_workbook

GRADE_LABELS = dict(zip((70,67,64,61,58,55,52,49,46,43,40),
                       ("A+","A","B+","B","B-","C+","C","C-","D+","D","E")))
SUMMARY_WORDS = ("小计", "合计", "总计")


def text(value):
    if value is None:
        return ""
    if isinstance(value, (float, int)) and not isinstance(value, bool) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def numeric(value, *, ratio=False, integer=False):
    if value is None or text(value) in ("", "—", "-", "缺考"):
        return None
    if isinstance(value, bool):
        raise ValueError("布尔值不是有效数字")
    s = text(value)
    percent = s.endswith("%")
    try:
        n = float(s.removesuffix("%"))
    except ValueError:
        raise ValueError(f"无效数字：{s}") from None
    if not math.isfinite(n) or n < 0:
        raise ValueError("数值必须非负且有限")
    if ratio:
        n = n / 100 if percent or n > 1 else n
        if n > 1:
            raise ValueError("比率/百分位必须介于 0 与 100%")
    elif percent:
        raise ValueError("此列不能填写百分比")
    if integer and not n.is_integer():
        raise ValueError("人数/排名必须为整数")
    return int(n) if integer else round(n, 10)


def parse(data: bytes, filename: str, grade: int, subject: str):
    wb = load_workbook(BytesIO(data), data_only=True)
    if len(wb.worksheets) != 1:
        raise ValueError("请按单张工作表上传；多工作表文件需拆分后核对")
    ws = wb.worksheets[0]
    result = {"filename": filename, "kind": "unknown", "rows": [], "classes": [],
              "aggregates": [], "excluded": [], "warnings": [], "sheet": ws.title}
    # 学校原表可在表头前带标题；仅扫描前十行，来源行号保留 Excel 原位。
    detail_header = next((r for r in range(1, min(10, ws.max_row) + 1)
                          if {"学号", "姓名"}.issubset(text(c.value) for c in ws[r])), None)
    if detail_header is not None:
        headers = [text(c.value) for c in ws[detail_header]]
        fields = {"学号": "alias", "姓名": "name", "分数": "score", "原始分数": "score",
                  "原始": "score", "原始分": "score", "等级": "grade_score", "等级分": "grade_score",
                  "等级分数": "grade_score", "学籍排名": "xueji_rank",
                  "年级百分位": "grade_percentile", "年级百分比": "grade_percentile"}
        cols = {fields[h]: i + 1 for i,h in enumerate(headers) if h in fields}
        if len(cols)!=len([h for h in headers if h in fields]):
            raise ValueError("明细存在重复字段列，请保留一列对应字段")
        if "score" not in cols:
            raise ValueError("缺少分数/原始分数列（补录时可以留空）")
        if (grade == 1 and any(h in headers for h in ("原始", "原始分", "原始分数"))) or (grade != 1 and "分数" in headers):
            raise ValueError("明细格式与所选年级不一致")
        result["kind"] = "student_scores"
        for r in range(detail_header + 1, ws.max_row + 1):
            if all(ws.cell(r,c).value is None for c in cols.values()):
                continue
            id_cell=ws.cell(r,cols["alias"])
            alias=text(id_cell.value)
            if isinstance(id_cell.value,(int,float)) and not isinstance(id_cell.value,bool) and re.fullmatch(r"0+",id_cell.number_format):
                alias=alias.zfill(len(id_cell.number_format))
            row = {"alias": alias,
                   "name": text(ws.cell(r, cols["name"]).value), "source_row": r}
            if not row["alias"] or not row["name"]:
                raise ValueError(f"第 {r} 行：学号和姓名均必填")
            for f,c in cols.items():
                if f in ("alias", "name"):
                    continue
                try:
                    row[f] = numeric(ws.cell(r,c).value, ratio=f=="grade_percentile", integer=f=="xueji_rank")
                except ValueError as exc:
                    raise ValueError(f"第 {r} 行 {headers[c-1]}：{exc}") from exc
            if row.get("xueji_rank") == 0:
                raise ValueError(f"第 {r} 行：排名须从 1 开始")
            if row.get("grade_score") is not None and row["grade_score"] not in GRADE_LABELS:
                raise ValueError(f"第 {r} 行：等级分不是已约定的 11 档")
            result["rows"].append(row)
        if not result["rows"]:
            raise ValueError("文件没有学生行")
        return result

    # 两行合并表头：只扩展实际合并区域，绝不填充教师空白。
    header_row = next((r for r in range(1, min(10,ws.max_row)+1)
                       if "班级" in [text(ws.cell(r,c).value) for c in range(1,ws.max_column+1)]), None)
    if header_row is None:
        raise ValueError("未识别教学单科成绩格式：前 10 行须有含学号、姓名的明细表头，或含班级的质量分析表头")
    def merged_value(r,c):
        v = ws.cell(r,c).value
        if v is None:
            for area in ws.merged_cells.ranges:
                if area.min_row <= r <= area.max_row and area.min_col <= c <= area.max_col:
                    return ws.cell(area.min_row,area.min_col).value
        return v
    tops = [text(merged_value(header_row,c)) for c in range(1,ws.max_column+1)]
    subs = [text(merged_value(header_row+1,c)) for c in range(1,ws.max_column+1)]
    if not any(t == "人数" for t in tops) or "应考" not in subs or "实考" not in subs:
        raise ValueError("质量表缺少应考/实考表头")
    is_g1 = "平均分" in tops and "原始" not in subs
    if is_g1 != (grade==1):
        raise ValueError("质量表格式与所选年级不一致")
    # 仅质量表按文件基本名判类；目录名不参与，文件名优先于可选的班级类型列。
    basename = re.sub(r"\s+", "", filename.replace("\\", "/").rsplit("/", 1)[-1])
    teaching_name, admin_name = "教学班" in basename, "行政班" in basename
    if teaching_name and admin_name:
        raise ValueError("文件名只能包含‘教学班’或‘行政班’其中一种，请改名后重新上传")
    if teaching_name:
        class_kind = "teaching"
    elif admin_name or "班级类型" in tops:
        class_kind = "administrative"
    else:
        raise ValueError("无法确定质量表类别，请在文件名中添加‘教学班’或‘行政班’后重新上传")
    result["class_kind"] = class_kind
    class_col = tops.index("班级")+1
    teacher_col = next((i+1 for i,h in enumerate(tops) if h in ("任课教师", "任教老师")), None)
    if teacher_col is None:
        raise ValueError("质量表缺少任课教师/任教老师列")
    type_col = tops.index("班级类型")+1 if "班级类型" in tops else None
    for h in tops:
        if "等级人数" in h and h != "等级人数及比率" and subject not in h:
            raise ValueError("质量表学科表头与所选学科不一致")
    metric_cols, band_cols = {}, {}
    last_band = None
    for c,(top,sub) in enumerate(zip(tops,subs),1):
        if sub in ("应考","实考"):
            metric_cols["expected_count" if sub=="应考" else "actual_count"] = c
        elif top == "平均分":
            metric_cols["grade_avg" if sub=="等级" else "raw_avg"] = c
        elif top in ("标准差","总分","最高分","最低分","语文","数学","英语","三大门"):
            metric_cols[{"标准差":"stddev", "总分":"score_sum", "最高分":"max_score", "最低分":"min_score",
                         "语文":"chinese_avg", "数学":"math_avg", "英语":"english_avg", "三大门":"main_three_avg"}[top]]=c
        if is_g1 and any(top.startswith(x) for x in ("优秀人数","良好人数","合格人数","低分人数")):
            band_cols.setdefault(top,{})["ratio" if sub=="比率" else "count"] = c
        elif not is_g1:
            if sub in GRADE_LABELS.values():
                last_band = sub
                band_cols.setdefault(sub,{})["count"] = c
            elif sub=="比率" and last_band:
                band_cols[last_band]["ratio"] = c
    if not is_g1 and set(band_cols) != set(GRADE_LABELS.values()):
        raise ValueError("质量表须完整包含 A+ 至 E 的 11 档表头")
    band_order=sorted(band_cols,key=lambda b:next(i for i,p in enumerate(("优秀","良好","合格","低分")) if b.startswith(p))) if is_g1 else list(GRADE_LABELS.values())
    current_type = ""
    coverage, group_coverage, coverage_counts = [], [], {}
    for r in range(header_row+2,ws.max_row+1):
        label = text(ws.cell(r,class_col).value)
        type_text = text(merged_value(r,type_col)) if type_col else ""
        teacher = text(ws.cell(r,teacher_col).value)
        summary = any(w in (label+type_text) for w in SUMMARY_WORDS)
        if teacher and not summary and not isinstance(ws.cell(r,teacher_col).value,str):
            raise ValueError(f"第 {r} 行：任课教师应填写姓名文字")
        if not label and not summary:
            continue
        if type_text and type_text != current_type and not summary:
            current_type = type_text
            group_coverage = []
        source = {"filename":filename,"sheet":ws.title,"row":r}
        # A:B 合并的“合计”行省掉教师列，数据整体比普通班级左移一列；
        # A:C 或 B:C 合并的小计则沿用普通列位。根据实际合并区域判断。
        shift = -1 if summary and any(a.min_row<=r<=a.max_row and a.min_col==1 and a.max_col==teacher_col-1 for a in ws.merged_cells.ranges) else 0
        metrics = {f:numeric(ws.cell(r,c+shift).value,integer=f.endswith("_count") and not summary) for f,c in metric_cols.items()}
        metrics["stddev_basis"] = "raw" if is_g1 else "unspecified"
        metrics["band_scheme"] = "score_range" if is_g1 else "grade11"
        metrics["bands"] = [{"label":b, **{f:numeric(ws.cell(r,c+shift).value,ratio=f=="ratio",integer=f=="count" and not summary) for f,c in band_cols[b].items()}} for b in band_order]
        expected,actual=metrics.get("expected_count"),metrics.get("actual_count")
        if expected is not None and actual is not None and actual>expected:
            raise ValueError(f"第 {r} 行：实考人数超过应考人数")
        if summary:
            covered=sorted(group_coverage if "小计" in label or ("合计" in type_text and not any(x in type_text for x in ("总计","闵中"))) else coverage)
            # 某些“合计”行实际上是班均人数（可为小数）。原样存储，
            # 只有人数确实等于覆盖班实考合计时才可作为总体汇总依据。
            verified=bool(covered) and actual is not None and all(coverage_counts[k] is not None for k in covered) and abs(sum(coverage_counts[k] for k in covered)-actual)<1e-8
            result["aggregates"].append({"label":label or type_text,"class_kind":class_kind,
                "coverage":covered,"coverage_verified":verified,
                "metrics":metrics,"source":source})
            result["excluded"].append({"class_label":label or type_text,"reason":"汇总行，不作为班级","row":r})
            continue
        key = class_kind+":"+label
        if key not in coverage:
            coverage.append(key)
        coverage_counts[key]=actual
        if key not in group_coverage:
            group_coverage.append(key)
        counts = [b.get("count") for b in metrics["bands"]]
        if not is_g1 and actual is not None and all(n is not None for n in counts) and sum(counts)!=actual:
            result["warnings"].append(f"第 {r} 行：等级人数合计与实考不符，等级总体均分不可合并")
        # 比率分母由学校定义，按官方值保留，不用实考或应考人数反算校验。
        result["classes"].append({"class_kind":class_kind,"class_label":label,
            "class_type":current_type or None,"teacher_name":teacher or None,"metrics":metrics,"source":source})
    result["kind"]="class_quality"
    if not result["classes"]:
        raise ValueError("质量表没有符合纳入条件的真实班级")
    return result
