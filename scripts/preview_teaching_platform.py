"""创建仅含合成数据的教学平台预览。必须显式提供全新数据目录。"""
import argparse
import os
from pathlib import Path
import sys


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", required=True)
    parser.add_argument("--files-dir", required=True)
    args = parser.parse_args()
    data = Path(args.data_dir).resolve()
    files = Path(args.files_dir).resolve()
    if data.exists():
        raise SystemExit("数据目录必须是新目录，拒绝覆盖")
    data.mkdir(parents=True, mode=0o700)
    files.mkdir(parents=True, exist_ok=True)
    os.environ["EXAM_TRACKER_DIR"] = str(data)
    os.environ["EXAM_TRACKER_BACKUP_DIR"] = str(data.parent / "backups")
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
    from datetime import date
    from io import BytesIO
    from openpyxl import Workbook
    from fastapi.testclient import TestClient
    from app.main import app
    from app.db.schema import ensure_app_schema
    from app.db.models import SessionLocal, Teacher
    from app.db.workspace_models import AcademicYear, TeachingClass
    from app.ingest.teaching_platform import GRADE_LABELS

    ensure_app_schema()
    db = SessionLocal()
    year = AcademicYear(name="2026-2027", start_date=date(2026, 9, 1), end_date=date(2027, 8, 31))
    db.add(year)
    db.add(Teacher(id=1, name="合成演示教师"))
    db.flush()
    classes = {}
    for grade in (1, 2, 3):
        tc = TeachingClass(academic_year_id=year.id, subject="物理", label=f"合成高{grade}班")
        db.add(tc)
        db.flush()
        classes[grade] = tc.id
    db.commit()

    def workbook(name, rows, merges=()):
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        for area in merges:
            wb.active.merge_cells(area)
        buf = BytesIO()
        wb.save(buf)
        path = files / name
        path.write_bytes(buf.getvalue())
        return name, buf.getvalue()

    def detail(grade, supplement=False):
        header = ["学号", "姓名", "分数", "学籍排名", "年级百分位"] if grade == 1 else ["学号", "姓名", "原始分数", "等级分数"]
        rows = [header]
        for i in range(1, 7):
            values = [f"SYN{grade}{i:03d}", f"合成学生{grade}{i}", 82-i*7]
            if grade == 1:
                values += [i*25, i*.05] if supplement else [None, None]
            else:
                values += [[67, 64, None, 55, 52, None][i-1] if supplement else None]
            rows.append(values)
        return workbook(f"合成高{grade}{'补录' if supplement else '明细'}.xlsx", rows)

    def quality(grade, kind):
        admin = kind == "administrative"
        prefix = ["班级类型"] if admin else []
        top = prefix + ["班级", "任课教师" if admin else "任教老师", "人数", None]
        sub = prefix + ["班级", "任课教师" if admin else "任教老师", "应考", "实考"]
        if grade == 1:
            top += ["总分", "平均分", "标准差", "最高分", "最低分"]
            sub += ["总分", "平均分", "标准差", "最高分", "最低分"]
            bands = ["优秀人数（100-85）", "良好人数（84-75）", "合格人数（74-60）", "低分人数（40分以下）"]
        else:
            top += ["平均分", None, "标准差"]
            sub += ["原始", "等级", "标准差"]
            bands = list(GRADE_LABELS.values())
        for b in bands:
            top += [b if grade == 1 else "等级人数及比率", None]
            sub += ["人数" if grade == 1 else b, "比率"]
        rows = [top, sub]
        labels = ["1", "2", "3"] if admin else [f"合成高{grade}班", f"合成高{grade}班a"]
        for i, label in enumerate(labels):
            row = (["合成平行班"] if admin else []) + [label, None if i == 1 else "合成教师", 32, 30]
            avg = [58.6, 49.8, 62.3][i]
            row += [avg*30, avg, 13.2+i, 94, 12] if grade == 1 else [avg, 55+i*3, 6.1+i]
            counts = [2, 4, 8, 5] if grade == 1 else [0, 1, 2, 3, 4, 5, 5, 4, 3, 2, 1]
            for count in counts:
                row += [count, round(count/30, 4)]
            rows.append(row)
        from openpyxl.utils import get_column_letter as col
        merges=[f'{col(len(prefix)+3)}1:{col(len(prefix)+4)}1']
        if grade!=1:
            merges.append(f'{col(len(prefix)+5)}1:{col(len(prefix)+6)}1')
        return workbook(f"合成高{grade}{'行政' if admin else '教学'}质量.xlsx", rows, merges)

    client = TestClient(app)
    results = []
    for grade in (1, 2, 3):
        initial = [detail(grade)]
        initial += [quality(grade, "administrative")] if grade == 1 else [quality(grade, "teaching"), quality(grade, "administrative")]
        base = {"academic_year_id": str(year.id), "grade": str(grade), "subject": "物理",
                "exam_name": f"合成高{grade}阶段测评", "exam_date": "2026-10-12", "teaching_class_id": str(classes[grade])}
        for operation, attachments in [("initial", initial), ("supplement", [detail(grade, True)])]:
            p = client.post("/api/v1/teaching/platform/imports/preview", data={**base,"operation":operation},
                files=[("files", (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for name, content in attachments])
            if p.status_code != 200 or p.json()["errors"]:
                raise RuntimeError(p.text)
            c = client.post("/api/v1/teaching/platform/imports/confirm", json={"token":p.json()["token"], "accept_roster_changes":True})
            if c.status_code != 200:
                raise RuntimeError(c.text)
        results.append({"grade":grade,"exam_id":c.json()["exam_id"]})
    db.close()
    print({"data_dir":str(data),"files_dir":str(files),"exams":results})


if __name__ == "__main__":
    main()
